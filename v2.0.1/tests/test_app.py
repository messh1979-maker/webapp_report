from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import inspect

from app import create_app
from extensions import db
from jalali import format_jalali_date, parse_jalali_date
from models import (
    Building, BuildingUsageType, City, EnergyProductionRecord, TelecomUsageRecord,
    User, UserModulePermission, UtilityBill, UtilitySubscription,
)


@pytest.fixture()
def app(tmp_path):
    application = create_app({
        "TESTING": True,
        "WTF_CSRF_ENABLED": False,
        "SECRET_KEY": "test-secret-not-for-production",
        "SQLALCHEMY_DATABASE_URI": f"sqlite:///{tmp_path / 'test.db'}",
    })
    with application.app_context():
        admin = User.query.filter_by(username="admin").first()
        admin.set_password("AdminStrong123")
        city = City.query.filter_by(code="T01").first()
        if not city:
            city = City(code="T01", name="شهر آزمون")
            db.session.add(city)
        db.session.commit()
    yield application


@pytest.fixture()
def client(app):
    client = app.test_client()
    response = client.post("/login", data={"username": "admin", "password": "AdminStrong123"})
    assert response.status_code == 302
    return client


def create_building():
    city = City.query.filter_by(code="T01").one()
    usage = BuildingUsageType.query.filter_by(code="office").one()
    building = Building(
        name="ساختمان آزمون",
        city=city,
        usage_type=usage,
        address="آدرس آزمون",
        staff_count=10,
        is_active=True,
    )
    db.session.add(building)
    db.session.flush()
    for utility, prefix in (("electricity", "E"), ("water", "W"), ("gas", "G")):
        db.session.add(UtilitySubscription(
            building=building, utility_type=utility,
            subscription_number=f"{prefix}-100", bill_identifier=f"B-{prefix}-100",
            installed_on=date(2026, 3, 21), activated_on=date(2026, 3, 22),
        ))
    db.session.commit()
    return building


def test_consumption_is_normalized_and_per_capita(app):
    with app.app_context():
        building = create_building()
        bill = UtilityBill(
            building=building,
            utility_subscription=building.subscription("electricity"),
            utility_type="electricity",
            period_type="periodic",
            year=1405,
            reading_from=date(2026, 3, 21),
            reading_to=date(2026, 4, 20),
            consumption=Decimal("125"),
            amount=Decimal("5000"),
        )
        db.session.add(bill)
        db.session.commit()
        assert bill.per_capita == Decimal("12.500")
        columns = {column["name"] for column in inspect(db.engine).get_columns("utility_bills")}
        assert "building_id" in columns
        assert "building_name" not in columns
        assert "subscription_number" not in columns
        assert "staff_count" not in columns


def test_periodic_rows_take_precedence_in_annual_report(app, client):
    with app.app_context():
        building = create_building()
        db.session.add_all([
            UtilityBill(building=building, utility_type="water", period_type="annual", year=1405,
                        reading_from=date(2026, 3, 21), reading_to=date(2027, 3, 20), consumption=999, amount=9999),
            UtilityBill(building=building, utility_type="water", period_type="periodic", year=1405,
                        reading_from=date(2026, 3, 21), reading_to=date(2026, 4, 20), consumption=30, amount=300),
            UtilityBill(building=building, utility_type="water", period_type="periodic", year=1405,
                        reading_from=date(2026, 4, 21), reading_to=date(2026, 5, 21), consumption=40, amount=400),
        ])
        db.session.commit()
    payload = client.get("/api/reports/data").get_json()
    assert payload["consumption"]["water"]["شهر آزمون"] == 70.0
    assert payload["amounts"]["water"]["شهر آزمون"] == 700.0


def test_expired_user_cannot_login(app):
    with app.app_context():
        city = City.query.filter_by(code="T01").one()
        user = User(username="expired", full_name="منقضی", role="user", assigned_city=city,
                    is_active_flag=True, active_until=datetime.now(UTC).replace(tzinfo=None) - timedelta(minutes=1))
        user.set_password("Expired12345")
        db.session.add(user)
        db.session.commit()
    response = app.test_client().post("/login", data={"username": "expired", "password": "Expired12345"})
    assert response.status_code == 401


def test_csrf_blocks_state_change(app, client):
    app.config["WTF_CSRF_ENABLED"] = True
    response = client.post("/logout", data={})
    assert response.status_code == 400


def test_security_headers(client):
    response = client.get("/")
    assert response.status_code == 200
    assert response.headers["X-Content-Type-Options"] == "nosniff"
    assert response.headers["X-Frame-Options"] == "DENY"
    assert "frame-ancestors 'none'" in response.headers["Content-Security-Policy"]


def test_building_excel_import_and_export(app, client):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["name", "city_code", "city_name", "usage_type", "address", "staff_count",
                  "electricity_subscription_number", "generator_type", "earth_resistance_ohm",
                  "earth_last_inspected_on", "is_active"])
    sheet.append(["ساختمان Import", "T01", "شهر آزمون", "اداری", "نشانی", 20, "E-100", "solar",
                  3.25, "۱۴۰۵/۰۶/۰۶", "بله"])
    subscriptions = workbook.create_sheet("اشتراک‌ها")
    subscriptions.append(["کد شهر", "نام ساختمان", "نوع انشعاب", "شماره اشتراک", "شناسه قبض",
                          "تاریخ نصب شمسی", "تاریخ به‌کارگیری شمسی", "تاریخ غیرفعال‌سازی شمسی"])
    subscriptions.append(["T01", "ساختمان Import", "برق", "E-100", "", "۱۴۰۵/۰۱/۰۱", "۱۴۰۵/۰۱/۰۲", ""])
    subscriptions.append(["T01", "ساختمان Import", "برق", "E-101", "B-101", "۱۴۰۵/۰۲/۰۱", "۱۴۰۵/۰۲/۰۲", ""])
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    response = client.post(
        "/buildings/import",
        data={"mode": "skip", "data_file": (stream, "buildings.xlsx")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    with app.app_context():
        building = Building.query.filter_by(name="ساختمان Import").one()
        assert len(building.utility_subscriptions) == 2
        assert building.subscription("electricity").subscription_number == "E-100"
        assert building.generator.generator_type == "solar"
        assert building.earth_system.resistance_ohm == Decimal("3.250")
        assert building.earth_system.last_inspected_on == date(2026, 8, 28)
    exported = client.get("/buildings/export")
    assert exported.status_code == 200
    assert exported.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_user_form_accepts_tehran_expiration(app, client):
    with app.app_context():
        city = City.query.filter_by(code="T01").one()
        city_id = city.id
    response = client.post("/admin/users/new", data={
        "username": "dated_user",
        "password": "DatedStrong123",
        "full_name": "کاربر تاریخ‌دار",
        "role": "user",
        "city_id": str(city_id),
        "is_active": "1",
        "active_from_date": "۱۴۰۵/۰۹/۱۰",
        "active_from_time": "09:00",
        "active_until_date": "۱۴۰۵/۱۰/۱۱",
        "active_until_time": "12:00",
    })
    assert response.status_code == 302
    with app.app_context():
        user = User.query.filter_by(username="dated_user").one()
        # 12:00 Asia/Tehran is persisted as 08:30 UTC.
        assert user.active_until == datetime(2027, 1, 1, 8, 30)


def test_user_excel_import_hashes_password(app, client):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["username", "password", "full_name", "phone", "role", "city_code", "is_active"])
    sheet.append(["import_user", "ImportStrong123", "کاربر Import", "", "user", "T01", "بله"])
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    response = client.post(
        "/admin/users/import",
        data={"excel_file": (stream, "users.xlsx")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    with app.app_context():
        user = User.query.filter_by(username="import_user").one()
        assert user.password_hash != "ImportStrong123"
        assert user.password_hash.startswith("scrypt:")
        assert user.check_password("ImportStrong123")


def test_jalali_conversion_and_bill_form_accepts_persian_money(app, client):
    assert parse_jalali_date("۱۴۰۵/۰۶/۰۶") == date(2026, 8, 28)
    assert format_jalali_date(date(2026, 8, 28)) == "۱۴۰۵/۰۶/۰۶"
    with app.app_context():
        building = create_building()
        building_id = building.id
        subscription_id = building.subscription("electricity").id
    response = client.post("/consumptions/new", data={
        "building_id": str(building_id),
        "utility_subscription_id": str(subscription_id),
        "utility_type": "electricity",
        "period_type": "periodic",
        "year": "۱۴۰۵",
        "reading_from": "۱۴۰۵/۰۱/۰۱",
        "reading_to": "۱۴۰۵/۰۲/۰۱",
        "consumption": "1250.5",
        "amount": "۱٬۲۳۴٬۵۶۷",
    })
    assert response.status_code == 302
    with app.app_context():
        bill = UtilityBill.query.filter_by(building_id=building_id).one()
        assert bill.reading_from == date(2026, 3, 21)
        assert bill.amount == Decimal("1234567.00")


def test_separate_excel_and_dynamic_reports(app, client):
    with app.app_context():
        building = create_building()
        db.session.add(UtilityBill(
            building=building, utility_subscription=building.subscription("water"),
            utility_type="water", period_type="annual", year=1405,
            reading_from=date(2026, 3, 21), reading_to=date(2027, 3, 20),
            consumption=Decimal("123.7"), amount=Decimal("9876543"),
        ))
        db.session.commit()
    response = client.get("/reports/export/water")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data), data_only=True)
    assert workbook.sheetnames == ["جزئیات دوره‌ها", "خلاصه سالیانه"]
    summary = workbook["خلاصه سالیانه"]
    assert summary.cell(2, 10).value == 124
    assert summary.cell(2, 12).value == 9876543
    assert summary.cell(2, 12).number_format == "#,##0"

    query = [
        ("dataset", "utility"), ("utility", "water"),
        ("field", "city_name"), ("field", "building_name"),
        ("field", "period_from"), ("field", "amount"),
    ]
    preview = client.get("/reports/dynamic", query_string=query)
    assert preview.status_code == 200
    assert "ساختمان آزمون".encode("utf-8") in preview.data
    exported = client.get("/reports/dynamic/export", query_string=query)
    assert exported.status_code == 200
    dynamic_book = load_workbook(BytesIO(exported.data), data_only=True)
    assert dynamic_book.active.max_row == 2


def test_login_displays_remaining_attempts(app):
    anonymous = app.test_client()
    first = anonymous.post("/login", data={"username": "admin", "password": "wrong-password"})
    assert first.status_code == 401
    assert b"<strong>4</strong>" in first.data
    second = anonymous.post("/login", data={"username": "admin", "password": "wrong-password"})
    assert second.status_code == 401
    assert b"<strong>3</strong>" in second.data


def test_module_permissions_and_telecom_rollout_flag(app):
    with app.app_context():
        city = City.query.filter_by(code="T01").one()
        user = User(username="scoped", full_name="کاربر محدود", role="user", assigned_city=city)
        user.set_password("ScopedStrong123")
        db.session.add(user)
        db.session.flush()
        db.session.add_all([
            UserModulePermission(user=user, module_key="gas", is_enabled=False),
            UserModulePermission(user=user, module_key="telecom", is_enabled=True),
        ])
        db.session.commit()
    scoped = app.test_client()
    assert scoped.post("/login", data={"username": "scoped", "password": "ScopedStrong123"}).status_code == 302
    assert scoped.get("/reports/export/gas").status_code == 403
    assert scoped.get("/telecom-usage").status_code == 403
    reports = scoped.get("/reports")
    assert reports.status_code == 200
    assert "گزارش گاز".encode("utf-8") not in reports.data
    assert "گزارش مخابرات".encode("utf-8") not in reports.data
    building_form = scoped.get("/buildings/new")
    assert "مخابرات، اینترنت و شبکه".encode("utf-8") not in building_form.data
    dynamic = scoped.get("/reports/dynamic?dataset=buildings")
    assert dynamic.status_code == 200
    assert b'f-telecom_type' not in dynamic.data

    app.config["TELECOM_USER_VISIBLE"] = True
    assert scoped.get("/telecom-usage").status_code == 200


def test_multiple_subscriptions_and_selected_subscription_dates(app, client):
    with app.app_context():
        building = create_building()
        second = UtilitySubscription(
            building=building, utility_type="electricity", subscription_number="E-200",
            bill_identifier="B-E-200", installed_on=date(2026, 4, 21), activated_on=date(2026, 4, 22),
        )
        db.session.add(second)
        db.session.commit()
        building_id, second_id = building.id, second.id
    api = client.get(f"/api/buildings/{building_id}").get_json()
    assert len(api["subscriptions"]["electricity"]) == 2
    selected = next(item for item in api["subscriptions"]["electricity"] if item["id"] == second_id)
    assert selected["installed_on"] == "۱۴۰۵/۰۲/۰۱"
    assert selected["activated_on"] == "۱۴۰۵/۰۲/۰۲"
    form = client.get("/consumptions/new")
    assert "تاریخ نصب انشعاب".encode("utf-8") in form.data
    response = client.post("/consumptions/new", data={
        "building_id": building_id, "utility_subscription_id": second_id,
        "utility_type": "electricity", "period_type": "periodic", "year": 1405,
        "reading_from": "۱۴۰۵/۰۲/۰۱", "reading_to": "۱۴۰۵/۰۳/۰۱",
        "consumption": "250", "amount": "۱۲٬۰۰۰",
    })
    assert response.status_code == 302
    with app.app_context():
        bill = UtilityBill.query.filter_by(utility_subscription_id=second_id).one()
        assert bill.building_id == building_id
    exported = load_workbook(BytesIO(client.get("/buildings/export").data), data_only=True)
    assert "اشتراک‌ها" in exported.sheetnames
    assert exported["اشتراک‌ها"].max_row == 5  # three defaults + second electricity


def test_production_is_grouped_by_generator_type(app, client):
    with app.app_context():
        building = create_building()
        building_id = building.id
    response = client.post("/production/new", data={
        "building_id": building_id, "year": 1405,
        "period_from": "۱۴۰۵/۰۱/۰۱", "period_to": "۱۴۰۵/۱۲/۲۹",
        "generator_type": "solar", "production_capacity_kwh": "500",
        "produced_kwh": "420", "notes": "آزمون خورشیدی",
    })
    assert response.status_code == 302
    with app.app_context():
        record = EnergyProductionRecord.query.one()
        assert record.generator_type == "solar"
        assert record.capacity_kwh == Decimal("500.000")
    workbook = load_workbook(BytesIO(client.get("/reports/export/production").data), data_only=True)
    assert workbook.sheetnames == ["جزئیات تولید", "تفکیک نوع ژنراتور"]
    summary = workbook["تفکیک نوع ژنراتور"]
    rows = {summary.cell(row, 1).value: summary.cell(row, 5).value for row in range(2, summary.max_row + 1)}
    assert rows["پنل خورشیدی"] == 420


def test_telecom_report_keeps_each_telecom_type_separate(app, client):
    with app.app_context():
        building = create_building()
        building_id = building.id
    for telecom_type, start, end in (
        ("e1", "۱۴۰۵/۰۱/۰۱", "۱۴۰۵/۰۲/۰۱"),
        ("sip", "۱۴۰۵/۰۲/۰۲", "۱۴۰۵/۰۳/۰۱"),
    ):
        response = client.post("/telecom-usage/new", data={
            "building_id": building_id, "year": 1405, "period_from": start, "period_to": end,
            "service_kind": "telecom", "telecom_type": telecom_type, "amount": "1000",
        })
        assert response.status_code == 302
    with app.app_context():
        assert {row.telecom_type for row in TelecomUsageRecord.query.all()} == {"e1", "sip"}
    workbook = load_workbook(BytesIO(client.get("/reports/export/telecom").data), data_only=True)
    values = {workbook.active.cell(row, 7).value for row in range(2, workbook.active.max_row + 1)}
    assert values == {"E1", "SIP"}


def test_city_per_capita_comparison_and_smart_building_dropdown(app, client):
    with app.app_context():
        first = create_building()
        first.address = "آدرس هوشمند"
        second_city = City(code="T02", name="شهر دوم")
        db.session.add(second_city)
        db.session.flush()
        usage = BuildingUsageType.query.filter_by(code="office").one()
        second = Building(name="ساختمان دوم", city=second_city, usage_type=usage, address="نشانی دوم", staff_count=20)
        db.session.add(second)
        db.session.flush()
        first_sub = first.subscription("water")
        second_sub = UtilitySubscription(building=second, utility_type="water", subscription_number="W-200")
        db.session.add(second_sub)
        db.session.add_all([
            UtilityBill(building=first, utility_subscription=first_sub, utility_type="water", period_type="periodic", year=1405, consumption=100, amount=1),
            UtilityBill(building=second, utility_subscription=second_sub, utility_type="water", period_type="periodic", year=1405, consumption=400, amount=1),
        ])
        db.session.commit()
        first_id = first.id
    payload = client.get("/api/reports/data?year=1405&period=periodic").get_json()
    rows = {row["city"]: row for row in payload["per_capita"]["water"]}
    assert rows["شهر آزمون"]["value"] == 10.0
    assert rows["شهر آزمون"]["others_average"] == 20.0
    assert rows["شهر آزمون"]["position"] == "below"
    assert rows["شهر دوم"]["position"] == "above"
    form = client.get("/buildings/new")
    assert f'value="{first_id}"'.encode() in form.data
    assert "آدرس هوشمند".encode("utf-8") in form.data
