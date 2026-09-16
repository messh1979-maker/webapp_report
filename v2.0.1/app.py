# -*- coding: utf-8 -*-
"""Normalized, secure energy monitoring application."""
import csv
import hashlib
import io
import json
import os
import re
import secrets
import time
from datetime import UTC, date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from functools import wraps
from pathlib import Path
from urllib.parse import urljoin, urlsplit
from zoneinfo import ZoneInfo

from flask import (
    Flask,
    abort,
    current_app,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from flask_login import current_user, login_required, login_user, logout_user
from flask_wtf.csrf import CSRFError
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import case, func, or_, select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import joinedload, selectinload
from werkzeug.middleware.proxy_fix import ProxyFix

from extensions import csrf, db, login_manager, migrate
from jalali import format_jalali_date, gregorian_to_jalali, normalize_digits, parse_jalali_date, to_persian_digits
from models import (
    Building,
    BuildingUsageType,
    City,
    EarthSystem,
    EnergyProductionRecord,
    LoginThrottle,
    NetworkService,
    PowerGenerator,
    TelecomService,
    TelecomUsageRecord,
    Tower,
    User,
    UserModulePermission,
    UtilityBill,
    UtilitySubscription,
    utcnow,
)

BASE_DIR = Path(__file__).resolve().parent
INSTANCE_DIR = BASE_DIR / "instance"

UTILITY_LABELS = {"electricity": "برق", "water": "آب", "gas": "گاز"}
UTILITY_UNITS = {"electricity": "kWh", "water": "m³", "gas": "m³"}
PERIOD_LABELS = {"annual": "سالیانه", "periodic": "دوره‌ای"}
GENERATOR_TYPES = {
    "diesel": "دیزل",
    "gas": "گازسوز",
    "hybrid": "هایبرید",
    "solar": "پنل خورشیدی",
    "wind": "بادی",
    "other": "سایر",
}
TELECOM_TYPES = {"e1": "E1", "analog": "آنالوگ", "sip_trunk": "SIP Trunk", "sip": "SIP"}
TOWER_TYPES = {"polygonal": "چندوجهی", "self_supporting": "خودایستا"}
SERVICE_KINDS = {"telecom": "مخابرات", "internet": "اینترنت", "mpls": "MPLS", "intranet": "Intranet"}
ROLES = {"admin": "ادمین", "admin_view": "ادمین ناظر", "user": "کاربر شهرستان", "view": "فقط مشاهده"}
MODULE_LABELS = {
    "buildings": "اطلاعات پایه ساختمان‌ها",
    "electricity": "مصارف و گزارش برق",
    "water": "مصارف و گزارش آب",
    "gas": "مصارف و گزارش گاز",
    "production": "فرم و گزارش تولید برق",
    "reports": "مرکز گزارش‌گیری",
    "dynamic_reports": "گزارش‌ساز پویا",
    "telecom": "فرم‌ها و گزارش مخابرات",
}
PAGE_SIZE = 50
MAX_LOGIN_FAILURES = 5
LOGIN_WINDOW_MINUTES = 15
LOGIN_LOCK_MINUTES = 15

BUILDING_COLUMNS = [
    ("name", "نام ساختمان"),
    ("city_code", "کد شهر"),
    ("city_name", "نام شهر"),
    ("usage_type", "کاربری ساختمان"),
    ("address", "آدرس ساختمان"),
    ("utm_zone", "زون UTM"),
    ("utm_easting", "UTM Easting"),
    ("utm_northing", "UTM Northing"),
    ("staff_count", "تعداد افراد"),
    ("electricity_subscription_number", "شماره اشتراک برق"),
    ("electricity_bill_identifier", "شناسه قبض برق"),
    ("electricity_installed_on", "تاریخ نصب برق (شمسی)"),
    ("electricity_activated_on", "تاریخ به‌کارگیری برق (شمسی)"),
    ("electricity_deactivated_on", "تاریخ غیرفعال‌سازی برق (شمسی)"),
    ("water_subscription_number", "شماره اشتراک آب"),
    ("water_bill_identifier", "شناسه قبض آب"),
    ("water_installed_on", "تاریخ نصب آب (شمسی)"),
    ("water_activated_on", "تاریخ به‌کارگیری آب (شمسی)"),
    ("water_deactivated_on", "تاریخ غیرفعال‌سازی آب (شمسی)"),
    ("gas_subscription_number", "شماره اشتراک گاز"),
    ("gas_bill_identifier", "شناسه قبض گاز"),
    ("gas_installed_on", "تاریخ نصب گاز (شمسی)"),
    ("gas_activated_on", "تاریخ به‌کارگیری گاز (شمسی)"),
    ("gas_deactivated_on", "تاریخ غیرفعال‌سازی گاز (شمسی)"),
    ("production_capacity_kwh", "ظرفیت تولید برق (kWh)"),
    ("generator_type", "نوع ژنراتور"),
    ("emergency_power_capacity_kw", "ظرفیت برق اضطراری (kW)"),
    ("phone_line_count", "تعداد خطوط تلفن"),
    ("telecom_type", "نوع سرویس مخابراتی"),
    ("telecom_provider", "شرکت سرویس‌دهنده مخابراتی"),
    ("telecom_case_number", "شماره پرونده مخابرات"),
    ("telecom_installed_on", "تاریخ نصب مخابرات (شمسی)"),
    ("internet_provider", "سرویس‌دهنده اینترنت"),
    ("internet_speed_mbps", "سرعت اینترنت (Mbps)"),
    ("internet_installed_on", "تاریخ نصب اینترنت (شمسی)"),
    ("mpls_installed_on", "تاریخ نصب MPLS (شمسی)"),
    ("intranet_installed_on", "تاریخ نصب Intranet (شمسی)"),
    ("tower_type", "نوع دکل"),
    ("tower_height_m", "ارتفاع دکل (m)"),
    ("has_earth", "دارای ارت"),
    ("earth_resistance_ohm", "مقاومت سیستم ارت (اهم)"),
    ("earth_last_inspected_on", "تاریخ شمسی آخرین اندازه‌گیری ارت"),
    ("is_active", "ساختمان فعال"),
    ("notes", "توضیحات"),
]


def _secret_key():
    """Persist a random local key; production should inject SECRET_KEY."""
    if os.environ.get("SECRET_KEY"):
        return os.environ["SECRET_KEY"]
    if os.environ.get("APP_ENV", "development").lower() == "production":
        raise RuntimeError("در محیط production تنظیم متغیر SECRET_KEY الزامی است.")
    INSTANCE_DIR.mkdir(parents=True, exist_ok=True)
    path = INSTANCE_DIR / ".secret_key"
    if not path.exists():
        path.write_text(secrets.token_urlsafe(48), encoding="ascii")
        try:
            path.chmod(0o600)
        except OSError:
            pass
    return path.read_text(encoding="ascii").strip()


def create_app(test_config=None):
    INSTANCE_DIR.mkdir(parents=True, exist_ok=True)
    app = Flask(__name__)
    # Trust forwarding headers only when exactly one controlled reverse proxy is used.
    if os.environ.get("TRUST_PROXY", "0") == "1":
        app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)
    database_url = os.environ.get("DATABASE_URL", f"sqlite:///{INSTANCE_DIR / 'app.db'}")
    engine_options = {"pool_pre_ping": True}
    if database_url.startswith("sqlite"):
        engine_options["connect_args"] = {"timeout": 15, "check_same_thread": False}

    app.config.update(
        SECRET_KEY=_secret_key(),
        SQLALCHEMY_DATABASE_URI=database_url,
        SQLALCHEMY_TRACK_MODIFICATIONS=False,
        SQLALCHEMY_ENGINE_OPTIONS=engine_options,
        MAX_CONTENT_LENGTH=10 * 1024 * 1024,
        WTF_CSRF_TIME_LIMIT=4 * 60 * 60,
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
        SESSION_COOKIE_SECURE=os.environ.get("COOKIE_SECURE", "0") == "1",
        PERMANENT_SESSION_LIFETIME=timedelta(minutes=int(os.environ.get("SESSION_MINUTES", "120"))),
        SESSION_REFRESH_EACH_REQUEST=False,
        APP_TIMEZONE=os.environ.get("APP_TIMEZONE", "Asia/Tehran"),
        # Telecom remains visible to administrators while this rollout flag is off.
        TELECOM_USER_VISIBLE=os.environ.get("TELECOM_USER_VISIBLE", "0") == "1",
        AUTO_CREATE_SCHEMA=os.environ.get("AUTO_CREATE_SCHEMA", "1") == "1",
    )
    if test_config:
        app.config.update(test_config)

    db.init_app(app)
    login_manager.init_app(app)
    csrf.init_app(app)
    migrate.init_app(app, db)
    app.jinja_env.filters["jalali_date"] = format_jalali_date
    app.jinja_env.filters["integer_number"] = _format_integer
    app.jinja_env.filters["rial"] = _format_rial

    @login_manager.user_loader
    def load_user(user_id):
        try:
            return db.session.get(User, int(user_id))
        except (TypeError, ValueError):
            return None

    register_hooks(app)
    register_routes(app)

    @app.context_processor
    def inject_globals():
        expiry = session.get("session_expires_at") if current_user.is_authenticated else None
        return {
            "utility_labels": UTILITY_LABELS,
            "utility_units": UTILITY_UNITS,
            "period_labels": PERIOD_LABELS,
            "generator_types": GENERATOR_TYPES,
            "telecom_types": TELECOM_TYPES,
            "tower_types": TOWER_TYPES,
            "service_kinds": SERVICE_KINDS,
            "roles": ROLES,
            "session_expires_at": expiry,
            "jalali_datetime_date": _jalali_datetime_date,
            "jalali_datetime_time": _jalali_datetime_time,
            "jalali_datetime_display": _jalali_datetime_display,
            "format_integer": _format_integer,
            "format_rial": _format_rial,
            "module_labels": MODULE_LABELS,
            "module_allowed": module_allowed,
            "telecom_visible": telecom_visible_for_user(),
            "allowed_utility_labels": {key: label for key, label in UTILITY_LABELS.items() if module_allowed(key)},
        }

    with app.app_context():
        # Run once before a multi-worker production server (see Dockerfile/README).
        # Local development keeps automatic initialization enabled.
        if app.config["AUTO_CREATE_SCHEMA"]:
            db.create_all()
            seed_reference_data()
            seed_admin()
            if database_url.startswith("sqlite"):
                db.session.execute(text("PRAGMA journal_mode=WAL"))
                db.session.execute(text("PRAGMA foreign_keys=ON"))
                db.session.commit()

    return app


def seed_reference_data():
    defaults = [("office", "اداری"), ("guest_house", "مهمان‌سرا"), ("warehouse", "انبار"), ("other", "سایر")]
    for code, name in defaults:
        if not BuildingUsageType.query.filter_by(code=code).first():
            db.session.add(BuildingUsageType(code=code, name=name))
    db.session.commit()


def seed_admin():
    if User.query.count() != 0:
        return
    city = City.query.filter_by(code="000").first()
    if not city:
        city = City(code="000", name="مرکز", is_active=True)
        db.session.add(city)
        db.session.flush()
    password = os.environ.get("INITIAL_ADMIN_PASSWORD") or secrets.token_urlsafe(14)
    admin = User(username="admin", full_name="مدیر سامانه", role="admin", is_active_flag=True)
    admin.set_password(password)
    db.session.add(admin)
    db.session.commit()
    print("=" * 68)
    print(f"مدیر اولیه ایجاد شد: admin / {password}")
    print("رمز را ذخیره و بلافاصله پس از ورود تغییر دهید.")
    print("=" * 68)


def register_hooks(app):
    @app.before_request
    def enforce_user_and_session_expiry():
        if not current_user.is_authenticated:
            return None
        now_ts = int(utcnow().replace(tzinfo=UTC).timestamp())
        expires_at = session.get("session_expires_at")
        if not current_user.is_active or (expires_at and now_ts >= int(expires_at)):
            logout_user()
            session.clear()
            flash("مهلت دسترسی یا نشست شما پایان یافته است.", "warning")
            return redirect(url_for("login"))
        return None

    @app.after_request
    def security_headers(response):
        response.headers.setdefault("X-Content-Type-Options", "nosniff")
        response.headers.setdefault("X-Frame-Options", "DENY")
        response.headers.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
        response.headers.setdefault("Permissions-Policy", "camera=(), microphone=(), geolocation=()")
        response.headers.setdefault(
            "Content-Security-Policy",
            "default-src 'self'; img-src 'self' data:; style-src 'self' 'unsafe-inline' https://cdn.jsdelivr.net https://cdnjs.cloudflare.com; "
            "font-src 'self' https://cdn.jsdelivr.net https://cdnjs.cloudflare.com; script-src 'self' 'unsafe-inline' https://cdn.jsdelivr.net; "
            "object-src 'none'; base-uri 'self'; frame-ancestors 'none'; form-action 'self'",
        )
        if request.endpoint not in {"static"}:
            response.headers.setdefault("Cache-Control", "no-store")
        return response

    @app.errorhandler(CSRFError)
    def handle_csrf(error):
        return render_template("error.html", code=400, message="درخواست نامعتبر یا منقضی شده است. صفحه را تازه‌سازی کنید."), 400

    @app.errorhandler(403)
    def forbidden(_error):
        return render_template("error.html", code=403, message="شما اجازه دسترسی به این بخش را ندارید."), 403

    @app.errorhandler(404)
    def not_found(_error):
        return render_template("error.html", code=404, message="صفحه مورد نظر یافت نشد."), 404

    @app.errorhandler(413)
    def too_large(_error):
        return render_template("error.html", code=413, message="حجم فایل بیش از ۱۰ مگابایت است."), 413


def admin_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not current_user.is_authenticated or not current_user.is_admin:
            abort(403)
        return fn(*args, **kwargs)
    return wrapped


def write_access_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not current_user.is_authenticated or not current_user.can_write:
            abort(403)
        return fn(*args, **kwargs)
    return wrapped


def telecom_visible_for_user():
    if not current_user.is_authenticated:
        return False
    return bool(current_user.is_admin or current_app.config.get("TELECOM_USER_VISIBLE", False))


def module_allowed(module_key):
    if not current_user.is_authenticated or module_key not in MODULE_LABELS:
        return False
    if current_user.is_admin:
        return True
    if module_key == "telecom" and not telecom_visible_for_user():
        return False
    return current_user.has_module(module_key)


def module_access_required(module_key):
    def decorator(fn):
        @wraps(fn)
        def wrapped(*args, **kwargs):
            if not module_allowed(module_key):
                abort(403)
            return fn(*args, **kwargs)
        return wrapped
    return decorator


def allowed_utilities():
    return tuple(key for key in UTILITY_LABELS if module_allowed(key))


def allowed_dynamic_datasets():
    datasets = []
    if module_allowed("buildings"):
        datasets.append("buildings")
    if allowed_utilities():
        datasets.append("utility")
    if module_allowed("production"):
        datasets.append("production")
    if module_allowed("telecom"):
        datasets.append("telecom")
    return tuple(datasets)


def visible_buildings_query():
    query = Building.query
    if not current_user.can_view_all_cities:
        query = query.filter(Building.city_id == current_user.city_id)
    return query


def visible_bills_query():
    query = UtilityBill.query.join(Building)
    if not current_user.can_view_all_cities:
        query = query.filter(Building.city_id == current_user.city_id)
    return query


def visible_production_query():
    query = EnergyProductionRecord.query.join(Building)
    if not current_user.can_view_all_cities:
        query = query.filter(Building.city_id == current_user.city_id)
    return query


def visible_telecom_query():
    query = TelecomUsageRecord.query.join(Building)
    if not current_user.can_view_all_cities:
        query = query.filter(Building.city_id == current_user.city_id)
    return query


def get_visible_building(building_id, write=False):
    building = db.session.get(Building, building_id)
    if not building:
        abort(404)
    if not current_user.can_view_all_cities and building.city_id != current_user.city_id:
        abort(403)
    if write and not current_user.can_write:
        abort(403)
    return building


def _safe_next(target):
    if not target:
        return None
    host = urlsplit(request.host_url)
    candidate = urlsplit(urljoin(request.host_url, target))
    return target if candidate.scheme in ("http", "https") and candidate.netloc == host.netloc else None


def _login_key(username):
    # ProxyFix (when explicitly enabled) has already replaced remote_addr safely.
    ip = request.remote_addr or "unknown"
    value = f"{username.casefold()}|{ip}|{current_app_secret_fingerprint()}"
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def current_app_secret_fingerprint():
    # Prevent precomputed identifiers without storing the secret or raw IP.
    return hashlib.sha256(current_app.config["SECRET_KEY"].encode("utf-8")).hexdigest()[:16]


def _register_failed_login(key):
    now = utcnow()
    throttle = LoginThrottle.query.filter_by(identifier_hash=key).first()
    if not throttle:
        throttle = LoginThrottle(identifier_hash=key, failure_count=0, window_started_at=now)
        db.session.add(throttle)
    if now - throttle.window_started_at > timedelta(minutes=LOGIN_WINDOW_MINUTES):
        throttle.failure_count = 0
        throttle.window_started_at = now
        throttle.locked_until = None
    throttle.failure_count += 1
    if throttle.failure_count >= MAX_LOGIN_FAILURES:
        throttle.locked_until = now + timedelta(minutes=LOGIN_LOCK_MINUTES)
    db.session.commit()
    time.sleep(min(1.25, 0.2 * throttle.failure_count))
    return throttle


def _parse_date(value, field_name="تاریخ", required=False):
    try:
        return parse_jalali_date(value, required=required)
    except ValueError as exc:
        if not value and required:
            raise ValueError(f"{field_name} الزامی است.") from exc
        raise ValueError(f"{field_name}: {exc}") from exc


def _app_timezone():
    try:
        return ZoneInfo(current_app.config.get("APP_TIMEZONE", "Asia/Tehran"))
    except Exception:
        return timezone(timedelta(hours=3, minutes=30), name="Asia/Tehran")


def _as_local_datetime(value):
    if not value:
        return None
    return value.replace(tzinfo=UTC).astimezone(_app_timezone())


def _jalali_datetime_date(value):
    local = _as_local_datetime(value)
    return format_jalali_date(local.date()) if local else ""


def _jalali_datetime_time(value):
    local = _as_local_datetime(value)
    return local.strftime("%H:%M") if local else ""


def _jalali_datetime_display(value):
    local = _as_local_datetime(value)
    return f"{format_jalali_date(local.date())}، {to_persian_digits(local.strftime('%H:%M'))}" if local else ""


def _parse_jalali_datetime(date_value, time_value, field_name, default_time="00:00"):
    if not date_value:
        return None
    parsed_date = _parse_date(date_value, field_name, required=True)
    raw_time = normalize_digits(time_value or default_time)
    match = re.fullmatch(r"(\d{1,2}):(\d{2})", raw_time)
    if not match:
        raise ValueError(f"ساعت {field_name} نامعتبر است.")
    hour, minute = map(int, match.groups())
    if hour > 23 or minute > 59:
        raise ValueError(f"ساعت {field_name} نامعتبر است.")
    local = datetime.combine(parsed_date, datetime.min.time()).replace(hour=hour, minute=minute, tzinfo=_app_timezone())
    return local.astimezone(UTC).replace(tzinfo=None)


def _parse_datetime_local(value, field_name):
    """Parse Jalali import values and UTC ISO values from secure JSON backups."""
    if value in (None, ""):
        return None
    try:
        if isinstance(value, datetime):
            parsed = value
        elif isinstance(value, date):
            parsed = datetime.combine(value, datetime.min.time())
        else:
            raw = normalize_digits(value)
            jalali_match = re.fullmatch(
                r"(1[2-6]\d{2}[\/\-.]\d{1,2}[\/\-.]\d{1,2})(?:[T\s،]+(\d{1,2}:\d{2}))?",
                raw,
            )
            if jalali_match:
                return _parse_jalali_datetime(jalali_match.group(1), jalali_match.group(2), field_name)
            parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=_app_timezone())
        return parsed.astimezone(UTC).replace(tzinfo=None)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{field_name} نامعتبر است؛ تاریخ را شمسی وارد کنید.") from exc


def _clean_numeric(value):
    return normalize_digits(value).replace(",", "").replace("٬", "").replace("،", "").replace(" ", "")


def _decimal(value, field_name, required=False, minimum=Decimal("0")):
    raw = _clean_numeric(value) if value is not None else ""
    if not raw:
        if required:
            raise ValueError(f"{field_name} الزامی است.")
        return None
    try:
        number = Decimal(raw)
    except InvalidOperation as exc:
        raise ValueError(f"{field_name} باید عدد معتبر باشد.") from exc
    if number < minimum:
        raise ValueError(f"{field_name} نمی‌تواند کمتر از {minimum} باشد.")
    return number


def _integer(value, field_name, required=False, minimum=0):
    raw = _clean_numeric(value) if value is not None else ""
    if not raw:
        if required:
            raise ValueError(f"{field_name} الزامی است.")
        return None
    try:
        decimal_value = Decimal(raw)
        if decimal_value != decimal_value.to_integral_value():
            raise InvalidOperation
        number = int(decimal_value)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{field_name} باید عدد صحیح باشد.") from exc
    if number < minimum:
        raise ValueError(f"{field_name} نمی‌تواند کمتر از {minimum} باشد.")
    return number


def _rounded_integer(value):
    if value in (None, ""):
        return None
    return int(Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _format_integer(value):
    number = _rounded_integer(value)
    if number is None:
        return "—"
    return to_persian_digits(f"{number:,}".replace(",", "٬"))


def _format_rial(value):
    formatted = _format_integer(value)
    return formatted if formatted == "—" else f"{formatted} ریال"


def _text(value, field_name, required=False, max_length=500):
    result = (str(value).strip() if value is not None else "")
    if required and not result:
        raise ValueError(f"{field_name} الزامی است.")
    if len(result) > max_length:
        raise ValueError(f"{field_name} حداکثر {max_length} نویسه مجاز است.")
    return result or None


def _bool(value, default=False):
    if value is None:
        return default
    return str(value).strip().casefold() in {"1", "true", "yes", "on", "بله", "فعال", "دارد"}


def _excel_safe(value):
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        return _jalali_datetime_display(value)
    if isinstance(value, date):
        return format_jalali_date(value)
    return value


def _reporting_year():
    today = date.today()
    return gregorian_to_jalali(today.year, today.month, today.day)[0]


def _jalali_stamp():
    return format_jalali_date(date.today(), persian_digits=False).replace("/", "")


def _form_value(source, key):
    value = source.get(key)
    return "" if value is None else value


def _get_or_create_related(building, relationship_name, model):
    obj = getattr(building, relationship_name)
    if not obj:
        obj = model(building=building)
        db.session.add(obj)
    return obj


def _find_city(source):
    city_id = source.get("city_id")
    city = db.session.get(City, int(city_id)) if city_id and str(city_id).isdigit() else None
    if not city:
        code = _text(source.get("city_code"), "کد شهر", required=True, max_length=30)
        name = _text(source.get("city_name"), "نام شهر", required=True, max_length=100)
        city = City.query.filter(or_(City.code == code, City.name == name)).first()
        if not city:
            city = City(code=code, name=name)
            db.session.add(city)
            db.session.flush()
    if not current_user.can_view_all_cities and city.id != current_user.city_id:
        abort(403)
    return city


def _find_usage(source):
    usage_id = source.get("usage_type_id")
    usage = db.session.get(BuildingUsageType, int(usage_id)) if usage_id and str(usage_id).isdigit() else None
    if usage:
        return usage
    raw = _text(source.get("usage_type"), "کاربری ساختمان", required=True, max_length=100)
    usage = BuildingUsageType.query.filter(or_(BuildingUsageType.code == raw, BuildingUsageType.name == raw)).first()
    if not usage:
        usage = BuildingUsageType.query.filter_by(code="other").first()
    return usage


def _source_values(source, key):
    if hasattr(source, "getlist"):
        return list(source.getlist(key))
    return [source.get(key)]


def _apply_utility_subscriptions(building, source, utility):
    """Apply indexed repeated form rows while preserving referenced subscriptions."""
    existing = {str(item.id): item for item in building.subscriptions(utility) if item.id is not None}
    ids = _source_values(source, f"{utility}_subscription_id")
    numbers = _source_values(source, f"{utility}_subscription_number")
    bills = _source_values(source, f"{utility}_bill_identifier")
    installed = _source_values(source, f"{utility}_installed_on")
    activated = _source_values(source, f"{utility}_activated_on")
    deactivated = _source_values(source, f"{utility}_deactivated_on")
    length = max(map(len, (ids, numbers, bills, installed, activated, deactivated)))
    submitted_ids = set()
    seen_numbers = set()
    seen_bill_ids = set()

    def at(values, index):
        return values[index] if index < len(values) else None

    for index in range(length):
        sub_id = str(at(ids, index) or "").strip()
        raw_number = at(numbers, index)
        raw_bill = at(bills, index)
        row_has_data = any(str(value or "").strip() for value in (
            raw_number, raw_bill, at(installed, index), at(activated, index), at(deactivated, index)
        ))
        sub = existing.get(sub_id) if sub_id else None
        if not sub and not hasattr(source, "getlist"):
            sub = building.subscription(utility)
        if sub_id and not sub:
            raise ValueError("شناسه اشتراک انتخاب‌شده معتبر نیست.")
        if not sub and not row_has_data:
            continue
        if not sub:
            sub = UtilitySubscription(building=building, utility_type=utility)
            db.session.add(sub)
        elif sub.utility_type != utility:
            raise ValueError("نوع اشتراک با رکورد انتخاب‌شده سازگار نیست.")
        if sub_id:
            submitted_ids.add(sub_id)
        sub.subscription_number = _text(raw_number, "شماره اشتراک", max_length=100)
        sub.bill_identifier = _text(raw_bill, "شناسه قبض", max_length=100)
        sub.installed_on = _parse_date(at(installed, index), "تاریخ نصب")
        sub.activated_on = _parse_date(at(activated, index), "تاریخ به‌کارگیری")
        sub.deactivated_on = _parse_date(at(deactivated, index), "تاریخ غیرفعال‌سازی")
        if sub.activated_on and sub.deactivated_on and sub.deactivated_on < sub.activated_on:
            raise ValueError(f"تاریخ غیرفعال‌سازی {UTILITY_LABELS[utility]} قبل از تاریخ به‌کارگیری است.")
        number_key = (sub.subscription_number or "").casefold()
        bill_key = (sub.bill_identifier or "").casefold()
        if number_key and number_key in seen_numbers:
            raise ValueError(f"شماره اشتراک تکراری برای {UTILITY_LABELS[utility]} وارد شده است.")
        if bill_key and bill_key in seen_bill_ids:
            raise ValueError(f"شناسه قبض تکراری برای {UTILITY_LABELS[utility]} وارد شده است.")
        if number_key:
            seen_numbers.add(number_key)
        if bill_key:
            seen_bill_ids.add(bill_key)

    # Dictionary imports update the preferred subscription only and never delete
    # additional rows. Repeated browser fields explicitly represent the full list.
    if hasattr(source, "getlist"):
        for sub_id, sub in existing.items():
            if sub_id in submitted_ids:
                continue
            if UtilityBill.query.filter_by(utility_subscription_id=sub.id).first():
                raise ValueError("اشتراک دارای قبض قابل حذف نیست؛ ابتدا قبض‌های وابسته را منتقل یا حذف کنید.")
            db.session.delete(sub)


def apply_building_data(building, source):
    city = _find_city(source)
    usage = _find_usage(source)
    building.name = _text(source.get("name"), "نام ساختمان", required=True, max_length=200)
    building.city = city
    building.usage_type = usage
    building.address = _text(source.get("address"), "آدرس ساختمان", required=True, max_length=700)
    building.utm_zone = _text(source.get("utm_zone"), "زون UTM", max_length=10)
    building.utm_easting = _decimal(source.get("utm_easting"), "UTM Easting")
    building.utm_northing = _decimal(source.get("utm_northing"), "UTM Northing")
    building.staff_count = _integer(source.get("staff_count"), "تعداد افراد", required=True) or 0
    if module_allowed("production"):
        building.emergency_power_capacity_kw = _decimal(source.get("emergency_power_capacity_kw"), "ظرفیت برق اضطراری")
    if module_allowed("telecom"):
        building.phone_line_count = _integer(source.get("phone_line_count"), "تعداد خطوط تلفن") or 0
    building.notes = _text(source.get("notes"), "توضیحات", max_length=4000)
    building.is_active = _bool(source.get("is_active"), default=True)

    for utility in UTILITY_LABELS:
        if module_allowed(utility):
            _apply_utility_subscriptions(building, source, utility)

    if module_allowed("production"):
        generator = _get_or_create_related(building, "generator", PowerGenerator)
        generator_type = source.get("generator_type") or "other"
        reverse_generators = {v: k for k, v in GENERATOR_TYPES.items()}
        generator.generator_type = reverse_generators.get(str(generator_type), str(generator_type))
        if generator.generator_type not in GENERATOR_TYPES:
            generator.generator_type = "other"
        generator.production_capacity_kwh = _decimal(source.get("production_capacity_kwh"), "ظرفیت تولید برق")

    if module_allowed("telecom"):
        telecom = _get_or_create_related(building, "telecom_service", TelecomService)
        telecom_type = source.get("telecom_type") or "analog"
        reverse_telecom = {v.casefold(): k for k, v in TELECOM_TYPES.items()}
        telecom.service_type = reverse_telecom.get(str(telecom_type).casefold(), str(telecom_type))
        if telecom.service_type not in TELECOM_TYPES:
            telecom.service_type = "analog"
        telecom.provider = _text(source.get("telecom_provider"), "سرویس‌دهنده مخابرات", max_length=200)
        telecom.case_number = _text(source.get("telecom_case_number"), "شماره پرونده مخابرات", max_length=100)
        telecom.installed_on = _parse_date(source.get("telecom_installed_on"), "تاریخ نصب مخابرات (شمسی)")

        networks = {n.service_kind: n for n in building.network_services}
        for kind in ("internet", "mpls", "intranet"):
            service = networks.get(kind)
            if not service:
                service = NetworkService(building=building, service_kind=kind)
                db.session.add(service)
            service.installed_on = _parse_date(source.get(f"{kind}_installed_on"), f"تاریخ نصب {kind}")
            if kind == "internet":
                service.provider = _text(source.get("internet_provider"), "سرویس‌دهنده اینترنت", max_length=200)
                service.speed_mbps = _decimal(source.get("internet_speed_mbps"), "سرعت اینترنت")

        tower = _get_or_create_related(building, "tower", Tower)
        tower_type = source.get("tower_type") or "self_supporting"
        reverse_towers = {v: k for k, v in TOWER_TYPES.items()}
        tower.tower_type = reverse_towers.get(str(tower_type), str(tower_type))
        if tower.tower_type not in TOWER_TYPES:
            tower.tower_type = "self_supporting"
        tower.height_m = _decimal(source.get("tower_height_m"), "ارتفاع دکل")

    earth = _get_or_create_related(building, "earth_system", EarthSystem)
    earth.has_earth = _bool(source.get("has_earth"), default=False)
    earth.resistance_ohm = _decimal(source.get("earth_resistance_ohm"), "مقاومت سیستم ارت")
    earth.last_inspected_on = _parse_date(source.get("earth_last_inspected_on"), "تاریخ اندازه‌گیری ارت")


def building_to_flat(building):
    data = {
        "name": building.name,
        "city_code": building.city.code,
        "city_name": building.city.name,
        "usage_type": building.usage_type.name,
        "address": building.address,
        "utm_zone": building.utm_zone,
        "utm_easting": building.utm_easting,
        "utm_northing": building.utm_northing,
        "staff_count": building.staff_count,
        "emergency_power_capacity_kw": building.emergency_power_capacity_kw,
        "phone_line_count": building.phone_line_count,
        "notes": building.notes,
        "is_active": "بله" if building.is_active else "خیر",
    }
    for utility in UTILITY_LABELS:
        sub = building.subscription(utility)
        for field in ("subscription_number", "bill_identifier", "installed_on", "activated_on", "deactivated_on"):
            data[f"{utility}_{field}"] = getattr(sub, field, None) if sub else None
    data["generator_type"] = GENERATOR_TYPES.get(building.generator.generator_type, "") if building.generator else None
    data["generator_type_code"] = building.generator.generator_type if building.generator else None
    data["production_capacity_kwh"] = building.generator.production_capacity_kwh if building.generator else None
    telecom = building.telecom_service
    data.update(
        telecom_type=TELECOM_TYPES.get(telecom.service_type, "") if telecom else None,
        telecom_provider=telecom.provider if telecom else None,
        telecom_case_number=telecom.case_number if telecom else None,
        telecom_installed_on=telecom.installed_on if telecom else None,
    )
    internet = building.network_service("internet")
    mpls = building.network_service("mpls")
    intranet = building.network_service("intranet")
    data.update(
        internet_provider=internet.provider if internet else None,
        internet_speed_mbps=internet.speed_mbps if internet else None,
        internet_installed_on=internet.installed_on if internet else None,
        mpls_installed_on=mpls.installed_on if mpls else None,
        intranet_installed_on=intranet.installed_on if intranet else None,
        tower_type=TOWER_TYPES.get(building.tower.tower_type, "") if building.tower else None,
        tower_height_m=building.tower.height_m if building.tower else None,
        has_earth="بله" if building.earth_system and building.earth_system.has_earth else "خیر",
        earth_resistance_ohm=building.earth_system.resistance_ohm if building.earth_system else None,
        earth_last_inspected_on=format_jalali_date(building.earth_system.last_inspected_on) if building.earth_system else None,
    )
    return data


def _style_worksheet(ws, integer_columns=(), currency_columns=()):
    ws.sheet_view.rightToLeft = True
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="343A40")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in set(integer_columns) | set(currency_columns):
        for column_cells in ws.iter_cols(min_col=column, max_col=column, min_row=2):
            for cell in column_cells:
                cell.number_format = "#,##0"
    for cells in ws.columns:
        max_length = min(55, max((len(str(cell.value or "")) for cell in cells), default=8) + 2)
        ws.column_dimensions[cells[0].column_letter].width = max(10, max_length)


def _workbook_response(headers, rows, title, filename, integer_columns=(), currency_columns=()):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for row in rows:
        ws.append([_excel_safe(value) for value in row])
    _style_worksheet(ws, integer_columns, currency_columns)
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return send_file(
        stream,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _multi_sheet_response(sheets, filename):
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        ws = wb.create_sheet(sheet["title"][:31])
        ws.append(sheet["headers"])
        for row in sheet["rows"]:
            ws.append([_excel_safe(value) for value in row])
        _style_worksheet(ws, sheet.get("integer_columns", ()), sheet.get("currency_columns", ()))
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return send_file(
        stream, as_attachment=True, download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _utility_annual_summaries(utility=None, year=None, building_id=None, utilities=None):
    aggregate = db.session.query(
        Building.id.label("building_id"), Building.name.label("building_name"),
        Building.address.label("address"), BuildingUsageType.name.label("usage_name"),
        City.name.label("city_name"), City.code.label("city_code"), Building.staff_count,
        UtilityBill.utility_type, UtilityBill.utility_subscription_id,
        UtilitySubscription.subscription_number, UtilitySubscription.bill_identifier, UtilityBill.year,
        func.sum(case((UtilityBill.period_type == "periodic", UtilityBill.consumption), else_=0)).label("periodic_consumption"),
        func.sum(case((UtilityBill.period_type == "periodic", UtilityBill.amount), else_=0)).label("periodic_amount"),
        func.sum(case((UtilityBill.period_type == "annual", UtilityBill.consumption), else_=0)).label("annual_consumption"),
        func.sum(case((UtilityBill.period_type == "annual", UtilityBill.amount), else_=0)).label("annual_amount"),
        func.sum(case((UtilityBill.period_type == "periodic", 1), else_=0)).label("periodic_count"),
    ).join(Building, UtilityBill.building_id == Building.id).join(City, Building.city_id == City.id).join(
        BuildingUsageType, Building.usage_type_id == BuildingUsageType.id
    ).outerjoin(UtilitySubscription, UtilityBill.utility_subscription_id == UtilitySubscription.id)
    if not current_user.can_view_all_cities:
        aggregate = aggregate.filter(Building.city_id == current_user.city_id)
    if utility in UTILITY_LABELS:
        aggregate = aggregate.filter(UtilityBill.utility_type == utility)
    elif utilities is not None:
        aggregate = aggregate.filter(UtilityBill.utility_type.in_(tuple(utilities)))
    if year:
        aggregate = aggregate.filter(UtilityBill.year == year)
    if building_id:
        aggregate = aggregate.filter(Building.id == building_id)
    grouped = aggregate.group_by(
        Building.id, Building.name, Building.address, BuildingUsageType.name,
        City.name, City.code, Building.staff_count, UtilityBill.utility_type,
        UtilityBill.utility_subscription_id, UtilitySubscription.subscription_number,
        UtilitySubscription.bill_identifier, UtilityBill.year
    ).order_by(UtilityBill.year.desc(), Building.name).all()
    result = []
    for row in grouped:
        use_periodic = row.periodic_count > 0
        consumption = row.periodic_consumption if use_periodic else row.annual_consumption
        amount = row.periodic_amount if use_periodic else row.annual_amount
        result.append({
            "building_id": row.building_id, "building_name": row.building_name, "address": row.address,
            "usage_name": row.usage_name, "city_name": row.city_name, "city_code": row.city_code,
            "staff_count": row.staff_count, "utility_type": row.utility_type,
            "utility_subscription_id": row.utility_subscription_id,
            "subscription_number": row.subscription_number, "bill_identifier": row.bill_identifier,
            "year": row.year, "consumption": consumption or 0, "amount": amount or 0,
            "source": "periodic" if use_periodic else "annual",
            "per_capita": (Decimal(consumption or 0) / Decimal(row.staff_count)).quantize(Decimal("0.001")) if row.staff_count else None,
        })
    return result


def _city_per_capita_comparison(period_type, year):
    if period_type not in PERIOD_LABELS:
        period_type = "periodic"
    aggregate = db.session.query(
        City.id.label("city_id"), City.name.label("city_name"), Building.id.label("building_id"),
        Building.staff_count, UtilityBill.utility_type, func.sum(UtilityBill.consumption).label("consumption"),
    ).join(Building, UtilityBill.building_id == Building.id).join(City, Building.city_id == City.id).filter(
        UtilityBill.period_type == period_type, UtilityBill.year == year,
        UtilityBill.utility_type.in_(allowed_utilities()), Building.staff_count > 0,
    ).group_by(City.id, City.name, Building.id, Building.staff_count, UtilityBill.utility_type).all()
    totals = {utility: {} for utility in allowed_utilities()}
    for row in aggregate:
        item = totals[row.utility_type].setdefault(
            row.city_id, {"city": row.city_name, "consumption": Decimal("0"), "staff": 0}
        )
        item["consumption"] += Decimal(row.consumption or 0)
        item["staff"] += int(row.staff_count or 0)
    result = {utility: [] for utility in allowed_utilities()}
    for utility, city_values in totals.items():
        capita = {
            city_id: (item["consumption"] / Decimal(item["staff"]))
            for city_id, item in city_values.items() if item["staff"]
        }
        visible_ids = list(capita)
        if not current_user.can_view_all_cities:
            visible_ids = [current_user.city_id] if current_user.city_id in capita else []
        for city_id in visible_ids:
            others = [value for other_id, value in capita.items() if other_id != city_id]
            others_average = sum(others, Decimal("0")) / Decimal(len(others)) if others else None
            value = capita[city_id]
            result[utility].append({
                "city": city_values[city_id]["city"],
                "value": float(value.quantize(Decimal("0.001"))),
                "others_average": float(others_average.quantize(Decimal("0.001"))) if others_average is not None else None,
                "delta": float((value - others_average).quantize(Decimal("0.001"))) if others_average is not None else None,
                "position": "above" if others_average is not None and value > others_average else (
                    "below" if others_average is not None and value < others_average else "equal"
                ),
            })
        result[utility].sort(key=lambda item: item["city"])
    return result


def _production_generator_summary():
    result = {
        key: {"generator_type": key, "label": label, "building_count": 0, "record_count": 0,
              "capacity": Decimal("0"), "produced": Decimal("0")}
        for key, label in GENERATOR_TYPES.items()
    }
    buildings = visible_buildings_query().options(selectinload(Building.generator)).all()
    for building in buildings:
        if not building.generator:
            continue
        key = building.generator.generator_type if building.generator.generator_type in GENERATOR_TYPES else "other"
        result[key]["building_count"] += 1
        result[key]["capacity"] += Decimal(building.generator.production_capacity_kwh or 0)
    records = visible_production_query().options(
        joinedload(EnergyProductionRecord.building).selectinload(Building.generator)
    ).all()
    for record in records:
        key = record.generator_type if record.generator_type in GENERATOR_TYPES else "other"
        result[key]["record_count"] += 1
        result[key]["produced"] += Decimal(record.produced_kwh or 0)
    return list(result.values())


def _building_of(item, dataset):
    return item if dataset == "buildings" else item.building


def _dynamic_field_specs(dataset):
    def building(item):
        return _building_of(item, dataset)

    common = {
        "city_name": ("نام شهر", lambda x: building(x).city.name, "text"),
        "city_code": ("کد شهر", lambda x: building(x).city.code, "text"),
        "building_name": ("نام ساختمان", lambda x: building(x).name, "text"),
        "usage_type": ("کاربری ساختمان", lambda x: building(x).usage_type.name, "text"),
        "address": ("آدرس ساختمان", lambda x: building(x).address, "text"),
        "staff_count": ("تعداد افراد", lambda x: building(x).staff_count, "integer"),
    }
    if dataset == "buildings":
        def sub(item, utility, field):
            record = item.subscription(utility)
            return getattr(record, field, None) if record else None
        def network(item, kind, field):
            record = item.network_service(kind)
            return getattr(record, field, None) if record else None
        return common | {
            "utm_zone": ("زون UTM", lambda x: x.utm_zone, "text"),
            "utm_easting": ("UTM Easting", lambda x: x.utm_easting, "decimal"),
            "utm_northing": ("UTM Northing", lambda x: x.utm_northing, "decimal"),
            "electricity_subscription": ("شماره اشتراک برق", lambda x: sub(x, "electricity", "subscription_number"), "text"),
            "electricity_bill_id": ("شناسه قبض برق", lambda x: sub(x, "electricity", "bill_identifier"), "text"),
            "electricity_installed": ("تاریخ نصب برق (شمسی)", lambda x: sub(x, "electricity", "installed_on"), "date"),
            "electricity_activated": ("تاریخ به‌کارگیری برق (شمسی)", lambda x: sub(x, "electricity", "activated_on"), "date"),
            "electricity_deactivated": ("تاریخ غیرفعال‌سازی برق (شمسی)", lambda x: sub(x, "electricity", "deactivated_on"), "date"),
            "water_subscription": ("شماره اشتراک آب", lambda x: sub(x, "water", "subscription_number"), "text"),
            "water_bill_id": ("شناسه قبض آب", lambda x: sub(x, "water", "bill_identifier"), "text"),
            "water_installed": ("تاریخ نصب آب (شمسی)", lambda x: sub(x, "water", "installed_on"), "date"),
            "water_activated": ("تاریخ به‌کارگیری آب (شمسی)", lambda x: sub(x, "water", "activated_on"), "date"),
            "water_deactivated": ("تاریخ غیرفعال‌سازی آب (شمسی)", lambda x: sub(x, "water", "deactivated_on"), "date"),
            "gas_subscription": ("شماره اشتراک گاز", lambda x: sub(x, "gas", "subscription_number"), "text"),
            "gas_bill_id": ("شناسه قبض گاز", lambda x: sub(x, "gas", "bill_identifier"), "text"),
            "gas_installed": ("تاریخ نصب گاز (شمسی)", lambda x: sub(x, "gas", "installed_on"), "date"),
            "gas_activated": ("تاریخ به‌کارگیری گاز (شمسی)", lambda x: sub(x, "gas", "activated_on"), "date"),
            "gas_deactivated": ("تاریخ غیرفعال‌سازی گاز (شمسی)", lambda x: sub(x, "gas", "deactivated_on"), "date"),
            "generator_type": ("نوع ژنراتور", lambda x: GENERATOR_TYPES.get(x.generator.generator_type, "") if x.generator else None, "text"),
            "production_capacity": ("ظرفیت تولید برق", lambda x: x.generator.production_capacity_kwh if x.generator else None, "integer"),
            "emergency_capacity": ("ظرفیت برق اضطراری", lambda x: x.emergency_power_capacity_kw, "integer"),
            "phone_lines": ("تعداد خطوط تلفن", lambda x: x.phone_line_count, "integer"),
            "telecom_type": ("نوع سرویس مخابرات", lambda x: TELECOM_TYPES.get(x.telecom_service.service_type, "") if x.telecom_service else None, "text"),
            "telecom_provider": ("سرویس‌دهنده مخابرات", lambda x: x.telecom_service.provider if x.telecom_service else None, "text"),
            "telecom_case": ("شماره پرونده مخابرات", lambda x: x.telecom_service.case_number if x.telecom_service else None, "text"),
            "telecom_installed": ("تاریخ نصب مخابرات (شمسی)", lambda x: x.telecom_service.installed_on if x.telecom_service else None, "date"),
            "internet_provider": ("سرویس‌دهنده اینترنت", lambda x: network(x, "internet", "provider"), "text"),
            "internet_speed": ("سرعت اینترنت Mbps", lambda x: network(x, "internet", "speed_mbps"), "decimal"),
            "internet_installed": ("تاریخ نصب اینترنت (شمسی)", lambda x: network(x, "internet", "installed_on"), "date"),
            "mpls_installed": ("تاریخ نصب MPLS (شمسی)", lambda x: network(x, "mpls", "installed_on"), "date"),
            "intranet_installed": ("تاریخ نصب Intranet (شمسی)", lambda x: network(x, "intranet", "installed_on"), "date"),
            "tower_type": ("نوع دکل", lambda x: TOWER_TYPES.get(x.tower.tower_type, "") if x.tower else None, "text"),
            "tower_height": ("ارتفاع دکل", lambda x: x.tower.height_m if x.tower else None, "decimal"),
            "has_earth": ("دارای سیستم ارت", lambda x: "بله" if x.earth_system and x.earth_system.has_earth else "خیر", "text"),
            "earth_resistance": ("مقاومت ارت (اهم)", lambda x: x.earth_system.resistance_ohm if x.earth_system else None, "decimal"),
            "earth_measured_on": ("تاریخ اندازه‌گیری ارت", lambda x: x.earth_system.last_inspected_on if x.earth_system else None, "date"),
            "notes": ("توضیحات ساختمان", lambda x: x.notes, "text"),
            "active": ("وضعیت ساختمان", lambda x: "فعال" if x.is_active else "غیرفعال", "text"),
            "created_by": ("ثبت‌کننده", lambda x: x.created_by.full_name if x.created_by else None, "text"),
            "created_at": ("زمان ثبت", lambda x: x.created_at, "datetime"),
            "updated_at": ("زمان آخرین ویرایش", lambda x: x.updated_at, "datetime"),
        }
    if dataset == "utility":
        return common | {
            "utility_type": ("نوع انشعاب", lambda x: UTILITY_LABELS[x.utility_type], "text"),
            "subscription_number": ("شماره اشتراک", lambda x: x.utility_subscription.subscription_number if x.utility_subscription else None, "text"),
            "bill_identifier": ("شناسه قبض", lambda x: x.utility_subscription.bill_identifier if x.utility_subscription else None, "text"),
            "subscription_installed": ("تاریخ نصب اشتراک", lambda x: x.utility_subscription.installed_on if x.utility_subscription else None, "date"),
            "subscription_activated": ("تاریخ به‌کارگیری اشتراک", lambda x: x.utility_subscription.activated_on if x.utility_subscription else None, "date"),
            "period_type": ("نوع دوره", lambda x: PERIOD_LABELS[x.period_type], "text"),
            "year": ("سال شمسی", lambda x: x.year, "integer_plain"),
            "period_from": ("از تاریخ", lambda x: x.reading_from, "date"),
            "period_to": ("تا تاریخ", lambda x: x.reading_to, "date"),
            "consumption": ("میزان مصرف", lambda x: x.consumption, "decimal"),
            "unit": ("واحد", lambda x: x.unit, "text"),
            "amount": ("مبلغ", lambda x: x.amount, "rial"),
            "per_capita": ("سرانه مصرف", lambda x: x.per_capita, "decimal"),
            "notes": ("توضیحات", lambda x: x.notes, "text"),
            "created_by": ("ثبت‌کننده", lambda x: x.created_by.full_name if x.created_by else None, "text"),
            "created_at": ("زمان ثبت", lambda x: x.created_at, "datetime"),
            "updated_at": ("زمان آخرین ویرایش", lambda x: x.updated_at, "datetime"),
        }
    if dataset == "production":
        return common | {
            "year": ("سال شمسی", lambda x: x.year, "integer_plain"),
            "period_from": ("از تاریخ", lambda x: x.period_from, "date"),
            "period_to": ("تا تاریخ", lambda x: x.period_to, "date"),
            "generator_type": ("نوع ژنراتور", lambda x: GENERATOR_TYPES.get(x.generator_type, "سایر"), "text"),
            "capacity": ("ظرفیت تولید", lambda x: x.capacity_kwh, "integer"),
            "produced_kwh": ("انرژی تولیدشده kWh", lambda x: x.produced_kwh, "integer"),
            "notes": ("توضیحات", lambda x: x.notes, "text"),
            "created_by": ("ثبت‌کننده", lambda x: x.created_by.full_name if x.created_by else None, "text"),
            "created_at": ("زمان ثبت", lambda x: x.created_at, "datetime"),
            "updated_at": ("زمان آخرین ویرایش", lambda x: x.updated_at, "datetime"),
        }
    return common | {
        "service_kind": ("نوع سرویس", lambda x: SERVICE_KINDS[x.service_kind], "text"),
        "year": ("سال شمسی", lambda x: x.year, "integer_plain"),
        "period_from": ("از تاریخ", lambda x: x.period_from, "date"),
        "period_to": ("تا تاریخ", lambda x: x.period_to, "date"),
        "telecom_type": ("نوع مخابرات", lambda x: TELECOM_TYPES.get(x.telecom_type, "") if x.telecom_type else None, "text"),
        "telecom_provider": ("سرویس‌دهنده مخابرات", lambda x: x.building.telecom_service.provider if x.building.telecom_service else None, "text"),
        "internet_provider": ("سرویس‌دهنده اینترنت", lambda x: (x.building.network_service("internet").provider if x.building.network_service("internet") else None), "text"),
        "amount": ("مبلغ", lambda x: x.amount, "rial"),
        "usage_description": ("شرح مصرف", lambda x: x.usage_description, "text"),
        "created_by": ("ثبت‌کننده", lambda x: x.created_by.full_name if x.created_by else None, "text"),
        "created_at": ("زمان ثبت", lambda x: x.created_at, "datetime"),
        "updated_at": ("زمان آخرین ویرایش", lambda x: x.updated_at, "datetime"),
    }


def _authorized_dynamic_field_specs(dataset):
    specs = _dynamic_field_specs(dataset)
    if dataset != "buildings":
        return specs
    hidden = set()
    for utility in UTILITY_LABELS:
        if not module_allowed(utility):
            hidden.update(key for key in specs if key.startswith(f"{utility}_"))
    if not module_allowed("production"):
        hidden.update({"generator_type", "production_capacity", "emergency_capacity"})
    if not module_allowed("telecom"):
        hidden.update({
            "phone_lines", "telecom_type", "telecom_provider", "telecom_case", "telecom_installed",
            "internet_provider", "internet_speed", "internet_installed", "mpls_installed",
            "intranet_installed", "tower_type", "tower_height",
        })
    return {key: value for key, value in specs.items() if key not in hidden}


def _dynamic_value(value, kind, excel=False):
    if value is None or value == "":
        return None if excel else "—"
    if kind == "date":
        return format_jalali_date(value)
    if kind == "datetime":
        return _jalali_datetime_display(value)
    if kind in ("integer", "rial"):
        return _rounded_integer(value) if excel else (_format_rial(value) if kind == "rial" else _format_integer(value))
    if kind == "integer_plain":
        return int(value) if excel else to_persian_digits(value)
    return _excel_safe(value) if excel else str(value)


def _dynamic_base_items(dataset, args, limit):
    city_id = args.get("city_id", type=int)
    building_id = args.get("building_id", type=int)
    year = args.get("year", type=int)
    utility = args.get("utility", "")
    service = args.get("service", "")
    date_from = _parse_date(args.get("date_from"), "از تاریخ") if args.get("date_from") else None
    date_to = _parse_date(args.get("date_to"), "تا تاریخ") if args.get("date_to") else None
    if date_from and date_to and date_to < date_from:
        raise ValueError("در گزارش پویا، تاریخ پایان قبل از تاریخ شروع است.")

    if dataset == "buildings":
        query = visible_buildings_query().options(
            selectinload(Building.utility_subscriptions), selectinload(Building.generator),
            selectinload(Building.telecom_service), selectinload(Building.network_services),
            selectinload(Building.tower), selectinload(Building.earth_system),
        )
        date_column = None
    elif dataset == "utility":
        available = allowed_utilities()
        if not available:
            abort(403)
        if utility and utility not in available:
            abort(403)
        query = visible_bills_query().options(
            joinedload(UtilityBill.building).selectinload(Building.utility_subscriptions),
            joinedload(UtilityBill.building).selectinload(Building.generator),
        ).filter(UtilityBill.utility_type.in_(available))
        if utility in UTILITY_LABELS:
            query = query.filter(UtilityBill.utility_type == utility)
        if year:
            query = query.filter(UtilityBill.year == year)
        date_column = UtilityBill.reading_to
    elif dataset == "production":
        query = visible_production_query().options(
            joinedload(EnergyProductionRecord.building).selectinload(Building.generator)
        )
        if year:
            query = query.filter(EnergyProductionRecord.year == year)
        date_column = EnergyProductionRecord.period_to
    else:
        query = visible_telecom_query().options(
            joinedload(TelecomUsageRecord.building).selectinload(Building.telecom_service),
            joinedload(TelecomUsageRecord.building).selectinload(Building.network_services),
        )
        if year:
            query = query.filter(TelecomUsageRecord.year == year)
        if service in SERVICE_KINDS:
            query = query.filter(TelecomUsageRecord.service_kind == service)
        date_column = TelecomUsageRecord.period_to
    if city_id:
        if not current_user.can_view_all_cities and city_id != current_user.city_id:
            abort(403)
        query = query.filter(Building.city_id == city_id)
    if building_id:
        get_visible_building(building_id)
        query = query.filter(Building.id == building_id)
    if date_column is not None:
        if date_from:
            query = query.filter(date_column >= date_from)
        if date_to:
            query = query.filter(date_column <= date_to)
    return query.limit(limit + 1).all()


def _dynamic_detail_result(dataset, args, limit=500, excel=False):
    specs = _authorized_dynamic_field_specs(dataset)
    defaults = {
        "buildings": ["city_name", "building_name", "usage_type", "address", "staff_count", "earth_resistance"],
        "utility": ["city_name", "building_name", "utility_type", "period_type", "year", "period_from", "period_to", "consumption", "amount", "per_capita"],
        "production": ["city_name", "building_name", "year", "period_from", "period_to", "capacity", "produced_kwh"],
        "telecom": ["city_name", "building_name", "service_kind", "telecom_type", "year", "period_from", "period_to", "amount", "usage_description"],
    }
    requested = args.getlist("field")
    selected = [key for key in requested if key in specs][:40] or defaults[dataset]
    items = _dynamic_base_items(dataset, args, limit)
    truncated = len(items) > limit
    items = items[:limit]
    sort_key = args.get("sort")
    if sort_key in specs:
        accessor = specs[sort_key][1]
        def sortable(item):
            value = accessor(item)
            return (value is None, value if value is not None else "")
        try:
            items.sort(key=sortable, reverse=args.get("dir") == "desc")
        except TypeError:
            items.sort(key=lambda item: str(accessor(item) or ""), reverse=args.get("dir") == "desc")
    columns = [{"key": key, "label": specs[key][0], "kind": specs[key][2]} for key in selected]
    rows = [[_dynamic_value(specs[key][1](item), specs[key][2], excel=excel) for key in selected] for item in items]
    return {"columns": columns, "rows": rows, "selected": selected, "truncated": truncated, "count": len(items)}


def register_routes(app):
    @app.get("/health")
    def health():
        db.session.execute(select(1)).scalar_one()
        return jsonify({"status": "ok"})

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if current_user.is_authenticated:
            return redirect(url_for("dashboard"))
        if request.method == "POST":
            username = str(request.form.get("username") or "").strip()[:80]
            password = request.form.get("password", "")
            key = _login_key(username)
            throttle = LoginThrottle.query.filter_by(identifier_hash=key).first()
            now = utcnow()
            if throttle and throttle.locked_until and throttle.locked_until > now:
                seconds = int((throttle.locked_until - now).total_seconds())
                flash(f"تلاش‌های ناموفق بیش از حد مجاز است. {seconds} ثانیه دیگر دوباره تلاش کنید.", "danger")
                return render_template(
                    "login.html", remaining_attempts=0, max_login_attempts=MAX_LOGIN_FAILURES,
                    login_username=username, locked_seconds=seconds,
                ), 429
            user = User.query.filter(func.lower(User.username) == username.casefold()).first()
            if not user or not user.check_password(password) or not user.is_active:
                throttle = _register_failed_login(key)
                if throttle.locked_until:
                    flash("حساب/نشانی شما به‌طور موقت قفل شد. ۱۵ دقیقه دیگر تلاش کنید.", "danger")
                else:
                    flash("نام کاربری یا رمز عبور نادرست است.", "danger")
                return render_template(
                    "login.html",
                    remaining_attempts=max(0, MAX_LOGIN_FAILURES - throttle.failure_count),
                    max_login_attempts=MAX_LOGIN_FAILURES,
                    login_username=username,
                ), 401

            LoginThrottle.query.filter_by(identifier_hash=key).delete()
            user.last_login_at = now
            db.session.commit()
            session.clear()
            session.permanent = True
            configured_expiry = now + app.config["PERMANENT_SESSION_LIFETIME"]
            effective_expiry = min(configured_expiry, user.active_until) if user.active_until else configured_expiry
            session["session_expires_at"] = int(effective_expiry.replace(tzinfo=UTC).timestamp())
            login_user(user, remember=False, fresh=True)
            flash(f"خوش آمدید، {user.full_name}", "success")
            return redirect(_safe_next(request.args.get("next")) or url_for("dashboard"))
        return render_template(
            "login.html", remaining_attempts=MAX_LOGIN_FAILURES,
            max_login_attempts=MAX_LOGIN_FAILURES, login_username="",
        )

    @app.post("/logout")
    @login_required
    def logout():
        logout_user()
        session.clear()
        flash("با موفقیت خارج شدید.", "info")
        return redirect(url_for("login"))

    @app.route("/")
    @login_required
    def dashboard():
        building_query = visible_buildings_query()
        bill_query = visible_bills_query()
        counts = {"buildings": building_query.count() if module_allowed("buildings") else 0}
        for utility in UTILITY_LABELS:
            counts[utility] = bill_query.filter(UtilityBill.utility_type == utility).count() if module_allowed(utility) else 0
        counts["production"] = visible_production_query().count() if module_allowed("production") else 0
        counts["telecom"] = visible_telecom_query().count() if module_allowed("telecom") else 0
        return render_template("dashboard.html", counts=counts)

    # ---------------- Buildings ----------------
    @app.route("/buildings")
    @login_required
    @module_access_required("buildings")
    def buildings():
        sort = request.args.get("sort", "name")
        direction = request.args.get("dir", "asc")
        columns = {
            "name": Building.name,
            "city": City.name,
            "usage": BuildingUsageType.name,
            "staff_count": Building.staff_count,
            "created_at": Building.created_at,
            "is_active": Building.is_active,
        }
        column = columns.get(sort, Building.name)
        direction = direction if direction in ("asc", "desc") else "asc"
        query = visible_buildings_query().join(City).join(BuildingUsageType)
        search = request.args.get("q", "").strip()
        if search:
            pattern = f"%{search[:100]}%"
            query = query.filter(or_(Building.name.ilike(pattern), Building.address.ilike(pattern), City.name.ilike(pattern)))
        query = query.options(
            selectinload(Building.utility_subscriptions),
            selectinload(Building.generator),
        ).order_by(column.asc() if direction == "asc" else column.desc())
        page = query.paginate(page=request.args.get("page", 1, type=int), per_page=PAGE_SIZE, error_out=False)
        return render_template("buildings.html", page=page, current_sort=sort, current_dir=direction, search=search)

    @app.route("/buildings/new", methods=["GET", "POST"])
    @login_required
    @module_access_required("buildings")
    @write_access_required
    def building_new():
        if request.method == "POST":
            try:
                existing_id = _integer(request.form.get("existing_building_id"), "ساختمان موجود", minimum=1)
                building = get_visible_building(existing_id, write=True) if existing_id else Building(created_by_id=current_user.id)
                # Validate/assign required scalar values before adding. Otherwise a
                # lookup query could autoflush an incomplete NOT NULL row.
                apply_building_data(building, request.form)
                if not existing_id:
                    db.session.add(building)
                db.session.commit()
                flash("اطلاعات پایه ساختمان به‌روزرسانی شد." if existing_id else "اطلاعات پایه ساختمان ثبت شد.", "success")
                return redirect(url_for("buildings"))
            except ValueError as exc:
                db.session.rollback()
                flash(str(exc), "danger")
            except Exception:
                db.session.rollback()
                app.logger.exception("building create failed")
                flash("ثبت ساختمان ناموفق بود؛ نام ساختمان در همان شهر نباید تکراری باشد.", "danger")
        cities = City.query.filter_by(is_active=True).order_by(City.name).all()
        usages = BuildingUsageType.query.order_by(BuildingUsageType.id).all()
        existing_buildings = visible_buildings_query().order_by(Building.name).all()
        return render_template(
            "building_form.html", building=None, cities=cities, usages=usages,
            existing_buildings=existing_buildings,
        )

    @app.route("/buildings/<int:building_id>/edit", methods=["GET", "POST"])
    @login_required
    @module_access_required("buildings")
    @write_access_required
    def building_edit(building_id):
        building = get_visible_building(building_id, write=True)
        if request.method == "POST":
            try:
                apply_building_data(building, request.form)
                db.session.commit()
                flash("اطلاعات ساختمان به‌روزرسانی شد.", "success")
                return redirect(url_for("buildings"))
            except ValueError as exc:
                db.session.rollback()
                flash(str(exc), "danger")
            except Exception:
                db.session.rollback()
                app.logger.exception("building update failed")
                flash("ویرایش ناموفق بود؛ داده‌های تکراری یا نامعتبر را بررسی کنید.", "danger")
        cities = City.query.filter_by(is_active=True).order_by(City.name).all()
        usages = BuildingUsageType.query.order_by(BuildingUsageType.id).all()
        return render_template("building_form.html", building=building, cities=cities, usages=usages)

    @app.post("/buildings/<int:building_id>/delete")
    @login_required
    @module_access_required("buildings")
    @write_access_required
    def building_delete(building_id):
        building = get_visible_building(building_id, write=True)
        db.session.delete(building)
        db.session.commit()
        flash("ساختمان و رکوردهای وابسته حذف شد.", "info")
        return redirect(url_for("buildings"))

    @app.get("/api/buildings/<int:building_id>")
    @login_required
    def building_api(building_id):
        if not (module_allowed("buildings") or allowed_utilities() or module_allowed("production") or module_allowed("telecom")):
            abort(403)
        building = get_visible_building(building_id)
        data = building_to_flat(building)
        data.update(
            id=building.id, city=building.city.name, city_id=building.city_id,
            usage=building.usage_type.name, usage_type_id=building.usage_type_id,
            is_active_value=bool(building.is_active),
        )
        for utility in UTILITY_LABELS:
            if not module_allowed(utility):
                for key in [item for item in data if item.startswith(f"{utility}_")]:
                    data.pop(key, None)
        if not module_allowed("production"):
            for key in ("generator_type", "generator_type_code", "production_capacity_kwh", "emergency_power_capacity_kw"):
                data.pop(key, None)
        if not module_allowed("telecom"):
            for key in list(data):
                if key.startswith(("telecom_", "internet_", "mpls_", "intranet_", "tower_")) or key == "phone_line_count":
                    data.pop(key, None)
        payload = {key: _excel_safe(value) for key, value in data.items()}
        payload["subscriptions"] = {
            utility: [{
                "id": sub.id,
                "subscription_number": sub.subscription_number or "",
                "bill_identifier": sub.bill_identifier or "",
                "installed_on": format_jalali_date(sub.installed_on),
                "activated_on": format_jalali_date(sub.activated_on),
                "deactivated_on": format_jalali_date(sub.deactivated_on),
                "active": sub.deactivated_on is None,
            } for sub in building.subscriptions(utility)]
            for utility in allowed_utilities()
        }
        return jsonify(payload)

    @app.get("/buildings/export")
    @login_required
    @module_access_required("buildings")
    def buildings_export():
        records = visible_buildings_query().options(
            selectinload(Building.utility_subscriptions),
            selectinload(Building.generator),
            selectinload(Building.telecom_service),
            selectinload(Building.network_services),
            selectinload(Building.tower),
            selectinload(Building.earth_system),
        ).order_by(Building.id).all()
        def column_visible(key):
            for utility in UTILITY_LABELS:
                if key.startswith(f"{utility}_") and not module_allowed(utility):
                    return False
            if key in {"production_capacity_kwh", "generator_type", "emergency_power_capacity_kw"} and not module_allowed("production"):
                return False
            if (key.startswith(("telecom_", "internet_", "mpls_", "intranet_", "tower_")) or key == "phone_line_count") and not module_allowed("telecom"):
                return False
            return True
        columns = [(key, label) for key, label in BUILDING_COLUMNS if column_visible(key)]
        rows = []
        subscription_rows = []
        for building in records:
            flat = building_to_flat(building)
            rows.append([flat.get(key) for key, _ in columns])
            for utility in allowed_utilities():
                for sub in building.subscriptions(utility):
                    subscription_rows.append([
                        building.city.code, building.city.name, building.name, UTILITY_LABELS[utility],
                        sub.subscription_number, sub.bill_identifier, sub.installed_on,
                        sub.activated_on, sub.deactivated_on,
                    ])
        return _multi_sheet_response([
            {"title": "ساختمان‌ها", "headers": [label for _, label in columns], "rows": rows},
            {"title": "اشتراک‌ها", "headers": [
                "کد شهر", "نام شهر", "نام ساختمان", "نوع انشعاب", "شماره اشتراک",
                "شناسه قبض", "تاریخ نصب شمسی", "تاریخ به‌کارگیری شمسی", "تاریخ غیرفعال‌سازی شمسی",
            ], "rows": subscription_rows},
        ], f"buildings_{_jalali_stamp()}.xlsx")

    @app.get("/buildings/import-template")
    @login_required
    @admin_required
    def buildings_import_template():
        example = {key: "" for key, _ in BUILDING_COLUMNS}
        example.update(name="ساختمان نمونه", city_code="091", city_name="مشهد", usage_type="اداری", address="نشانی نمونه", staff_count=10, earth_resistance_ohm=2.5, earth_last_inspected_on="۱۴۰۵/۰۶/۰۶", is_active="بله")
        return _multi_sheet_response([
            {"title": "ساختمان‌ها", "headers": [label for _, label in BUILDING_COLUMNS],
             "rows": [[example[key] for key, _ in BUILDING_COLUMNS]]},
            {"title": "اشتراک‌ها", "headers": [
                "کد شهر", "نام ساختمان", "نوع انشعاب", "شماره اشتراک", "شناسه قبض",
                "تاریخ نصب شمسی", "تاریخ به‌کارگیری شمسی", "تاریخ غیرفعال‌سازی شمسی",
            ], "rows": [
                ["091", "ساختمان نمونه", "برق", "E-100", "B-100", "۱۴۰۵/۰۱/۰۱", "۱۴۰۵/۰۱/۰۲", ""],
                ["091", "ساختمان نمونه", "برق", "E-101", "B-101", "۱۴۰۵/۰۲/۰۱", "۱۴۰۵/۰۲/۰۲", ""],
            ]},
        ], "building_import_template.xlsx")

    @app.route("/buildings/import", methods=["GET", "POST"])
    @login_required
    @admin_required
    def buildings_import():
        result = None
        if request.method == "POST":
            upload = request.files.get("data_file")
            if not upload or not upload.filename:
                flash("فایل Excel یا CSV را انتخاب کنید.", "danger")
                return redirect(url_for("buildings_import"))
            extension = Path(upload.filename).suffix.lower()
            try:
                subscription_raw_rows = []
                if extension in (".xlsx", ".xlsm"):
                    workbook = load_workbook(upload, read_only=True, data_only=True)
                    ws = workbook.active
                    raw_rows = list(ws.iter_rows(values_only=True))
                    if "اشتراک‌ها" in workbook.sheetnames:
                        subscription_raw_rows = list(workbook["اشتراک‌ها"].iter_rows(values_only=True))
                elif extension == ".csv":
                    text = upload.stream.read().decode("utf-8-sig")
                    raw_rows = list(csv.reader(io.StringIO(text)))
                else:
                    raise ValueError("فقط فایل xlsx، xlsm یا CSV مجاز است.")
                if not raw_rows:
                    raise ValueError("فایل خالی است.")
                labels_to_keys = {label: key for key, label in BUILDING_COLUMNS} | {key: key for key, _ in BUILDING_COLUMNS}
                headers = [labels_to_keys.get(str(v).strip()) if v is not None else None for v in raw_rows[0]]
                if "name" not in headers:
                    raise ValueError("ستون «نام ساختمان» در فایل وجود ندارد.")
                mode = request.form.get("mode", "skip")
                created, updated, skipped, errors = [], [], [], []
                for row_number, row in enumerate(raw_rows[1:], 2):
                    if not row or all(v is None or str(v).strip() == "" for v in row):
                        continue
                    mapping = {headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]}
                    try:
                        with db.session.begin_nested():
                            city = _find_city(mapping)
                            name = _text(mapping.get("name"), "نام ساختمان", required=True, max_length=200)
                            existing = Building.query.filter_by(city_id=city.id, name=name).first()
                            if existing and mode != "update":
                                skipped.append(f"ردیف {row_number}: {name}")
                                continue
                            building = existing or Building(created_by_id=current_user.id)
                            apply_building_data(building, mapping)
                            if not existing:
                                db.session.add(building)
                            db.session.flush()
                            (updated if existing else created).append(name)
                    except (ValueError, TypeError, IntegrityError) as exc:
                        errors.append(f"ردیف {row_number}: داده نامعتبر یا تکراری است ({exc.__class__.__name__}).")

                if subscription_raw_rows:
                    sub_header_map = {
                        "کد شهر": "city_code", "نام ساختمان": "building_name", "نوع انشعاب": "utility_type",
                        "شماره اشتراک": "subscription_number", "شناسه قبض": "bill_identifier",
                        "تاریخ نصب شمسی": "installed_on", "تاریخ به‌کارگیری شمسی": "activated_on",
                        "تاریخ غیرفعال‌سازی شمسی": "deactivated_on",
                    }
                    sub_headers = [sub_header_map.get(str(value).strip()) if value is not None else None for value in subscription_raw_rows[0]]
                    reverse_utilities = {label: key for key, label in UTILITY_LABELS.items()} | {key: key for key in UTILITY_LABELS}
                    for row_number, row in enumerate(subscription_raw_rows[1:10001], 2):
                        if not row or all(value is None or str(value).strip() == "" for value in row):
                            continue
                        mapping = {sub_headers[i]: row[i] for i in range(min(len(sub_headers), len(row))) if sub_headers[i]}
                        try:
                            city_code = _text(mapping.get("city_code"), "کد شهر", required=True, max_length=30)
                            building_name = _text(mapping.get("building_name"), "نام ساختمان", required=True, max_length=200)
                            utility = reverse_utilities.get(str(mapping.get("utility_type") or "").strip())
                            if utility not in UTILITY_LABELS:
                                raise ValueError("نوع انشعاب معتبر نیست.")
                            building = Building.query.join(City).filter(City.code == city_code, Building.name == building_name).first()
                            if not building:
                                raise ValueError("ساختمان متناظر یافت نشد.")
                            number = _text(mapping.get("subscription_number"), "شماره اشتراک", max_length=100)
                            bill_identifier = _text(mapping.get("bill_identifier"), "شناسه قبض", max_length=100)
                            duplicate = UtilitySubscription.query.filter_by(
                                building_id=building.id, utility_type=utility,
                                subscription_number=number, bill_identifier=bill_identifier,
                            ).first()
                            if duplicate:
                                sub = duplicate
                            else:
                                sub = UtilitySubscription(building=building, utility_type=utility)
                                db.session.add(sub)
                            sub.subscription_number = number
                            sub.bill_identifier = bill_identifier
                            sub.installed_on = _parse_date(mapping.get("installed_on"), "تاریخ نصب")
                            sub.activated_on = _parse_date(mapping.get("activated_on"), "تاریخ به‌کارگیری")
                            sub.deactivated_on = _parse_date(mapping.get("deactivated_on"), "تاریخ غیرفعال‌سازی")
                            if sub.activated_on and sub.deactivated_on and sub.deactivated_on < sub.activated_on:
                                raise ValueError("تاریخ غیرفعال‌سازی قبل از به‌کارگیری است.")
                        except (ValueError, TypeError, IntegrityError) as exc:
                            errors.append(f"Sheet اشتراک‌ها، ردیف {row_number}: {exc}")
                db.session.commit()
                result = {"created": created, "updated": updated, "skipped": skipped, "errors": errors}
                flash(f"Import پایان یافت: {len(created)} ایجاد و {len(updated)} به‌روزرسانی شد.", "success")
            except (ValueError, UnicodeDecodeError) as exc:
                db.session.rollback()
                flash(str(exc), "danger")
            except Exception:
                db.session.rollback()
                app.logger.exception("building import failed")
                flash("فایل قابل پردازش نیست یا ساختار آن نامعتبر است.", "danger")
        return render_template("building_import.html", result=result)

    # ---------------- Utility bills ----------------
    @app.get("/consumptions")
    @login_required
    def consumptions():
        available = allowed_utilities()
        if not available:
            abort(403)
        utility = request.args.get("utility", "")
        if utility and utility not in available:
            abort(403)
        year = request.args.get("year", type=int)
        sort = request.args.get("sort", "reading_to")
        direction = request.args.get("dir", "desc")
        columns = {
            "building": Building.name,
            "city": City.name,
            "utility": UtilityBill.utility_type,
            "period": UtilityBill.period_type,
            "year": UtilityBill.year,
            "reading_from": UtilityBill.reading_from,
            "reading_to": UtilityBill.reading_to,
            "consumption": UtilityBill.consumption,
            "amount": UtilityBill.amount,
        }
        column = columns.get(sort, UtilityBill.reading_to)
        direction = direction if direction in ("asc", "desc") else "desc"
        query = visible_bills_query().join(City, Building.city_id == City.id).filter(
            UtilityBill.utility_type.in_(available)
        )
        if utility in UTILITY_LABELS:
            query = query.filter(UtilityBill.utility_type == utility)
        else:
            utility = ""
        if year:
            query = query.filter(UtilityBill.year == year)
        query = query.order_by(column.asc() if direction == "asc" else column.desc(), UtilityBill.id.desc())
        page = query.paginate(page=request.args.get("page", 1, type=int), per_page=PAGE_SIZE, error_out=False)

        summaries = _utility_annual_summaries(utility=utility, year=year, utilities=available)
        return render_template("consumptions.html", page=page, summaries=summaries, current_utility=utility, current_year=year, current_sort=sort, current_dir=direction)

    @app.route("/consumptions/new", methods=["GET", "POST"])
    @login_required
    @write_access_required
    def consumption_new():
        bill = UtilityBill(created_by_id=current_user.id)
        return _bill_form(app, bill, is_new=True)

    @app.route("/consumptions/<int:bill_id>/edit", methods=["GET", "POST"])
    @login_required
    @write_access_required
    def consumption_edit(bill_id):
        bill = db.session.get(UtilityBill, bill_id)
        if not bill:
            abort(404)
        if not module_allowed(bill.utility_type):
            abort(403)
        get_visible_building(bill.building_id, write=True)
        return _bill_form(app, bill, is_new=False)

    @app.post("/consumptions/<int:bill_id>/delete")
    @login_required
    @write_access_required
    def consumption_delete(bill_id):
        bill = db.session.get(UtilityBill, bill_id)
        if not bill:
            abort(404)
        if not module_allowed(bill.utility_type):
            abort(403)
        get_visible_building(bill.building_id, write=True)
        db.session.delete(bill)
        db.session.commit()
        flash("رکورد قبض حذف شد.", "info")
        return redirect(url_for("consumptions"))

    @app.get("/consumptions/export")
    @login_required
    def consumptions_export():
        available = allowed_utilities()
        if not available:
            abort(403)
        records = visible_bills_query().filter(UtilityBill.utility_type.in_(available)).order_by(
            UtilityBill.year.desc(), UtilityBill.id
        ).all()
        headers = ["شهر", "ساختمان", "انشعاب", "شماره اشتراک", "شناسه قبض", "تاریخ نصب", "تاریخ به‌کارگیری", "نوع دوره", "سال شمسی", "از تاریخ شمسی", "تا تاریخ شمسی", "مصرف", "واحد", "مبلغ (ریال)", "سرانه"]
        rows = [[
            r.building.city.name, r.building.name, UTILITY_LABELS[r.utility_type],
            r.utility_subscription.subscription_number if r.utility_subscription else None,
            r.utility_subscription.bill_identifier if r.utility_subscription else None,
            r.utility_subscription.installed_on if r.utility_subscription else None,
            r.utility_subscription.activated_on if r.utility_subscription else None,
            PERIOD_LABELS[r.period_type], r.year, r.reading_from, r.reading_to,
            r.consumption, r.unit, _rounded_integer(r.amount), r.per_capita,
        ] for r in records]
        return _workbook_response(headers, rows, "مصارف", f"consumptions_{_jalali_stamp()}.xlsx", currency_columns=(14,))

    # ---------------- Energy production ----------------
    @app.get("/production")
    @login_required
    @module_access_required("production")
    def production_list():
        query = EnergyProductionRecord.query.join(Building)
        if not current_user.can_view_all_cities:
            query = query.filter(Building.city_id == current_user.city_id)
        sort = request.args.get("sort", "year")
        direction = request.args.get("dir", "desc")
        column = {"building": Building.name, "year": EnergyProductionRecord.year, "produced_kwh": EnergyProductionRecord.produced_kwh, "period_to": EnergyProductionRecord.period_to}.get(sort, EnergyProductionRecord.year)
        query = query.order_by(column.asc() if direction == "asc" else column.desc())
        page = query.paginate(page=request.args.get("page", 1, type=int), per_page=PAGE_SIZE, error_out=False)
        return render_template("service_records.html", page=page, kind="production", current_sort=sort, current_dir=direction)

    @app.route("/production/new", methods=["GET", "POST"])
    @login_required
    @module_access_required("production")
    @write_access_required
    def production_new():
        record = EnergyProductionRecord(created_by_id=current_user.id)
        return _service_form(app, record, "production", True)

    @app.route("/production/<int:record_id>/edit", methods=["GET", "POST"])
    @login_required
    @module_access_required("production")
    @write_access_required
    def production_edit(record_id):
        record = db.session.get(EnergyProductionRecord, record_id)
        if not record:
            abort(404)
        get_visible_building(record.building_id, write=True)
        return _service_form(app, record, "production", False)

    @app.post("/production/<int:record_id>/delete")
    @login_required
    @module_access_required("production")
    @write_access_required
    def production_delete(record_id):
        record = db.session.get(EnergyProductionRecord, record_id)
        if not record:
            abort(404)
        get_visible_building(record.building_id, write=True)
        db.session.delete(record)
        db.session.commit()
        return redirect(url_for("production_list"))

    # ---------------- Telecom usage ----------------
    @app.get("/telecom-usage")
    @login_required
    @module_access_required("telecom")
    def telecom_usage_list():
        query = TelecomUsageRecord.query.join(Building)
        if not current_user.can_view_all_cities:
            query = query.filter(Building.city_id == current_user.city_id)
        sort = request.args.get("sort", "year")
        direction = request.args.get("dir", "desc")
        column = {"building": Building.name, "year": TelecomUsageRecord.year, "amount": TelecomUsageRecord.amount, "service_kind": TelecomUsageRecord.service_kind, "period_to": TelecomUsageRecord.period_to}.get(sort, TelecomUsageRecord.year)
        query = query.order_by(column.asc() if direction == "asc" else column.desc())
        page = query.paginate(page=request.args.get("page", 1, type=int), per_page=PAGE_SIZE, error_out=False)
        return render_template("service_records.html", page=page, kind="telecom", current_sort=sort, current_dir=direction)

    @app.route("/telecom-usage/new", methods=["GET", "POST"])
    @login_required
    @module_access_required("telecom")
    @write_access_required
    def telecom_usage_new():
        record = TelecomUsageRecord(created_by_id=current_user.id)
        return _service_form(app, record, "telecom", True)

    @app.route("/telecom-usage/<int:record_id>/edit", methods=["GET", "POST"])
    @login_required
    @module_access_required("telecom")
    @write_access_required
    def telecom_usage_edit(record_id):
        record = db.session.get(TelecomUsageRecord, record_id)
        if not record:
            abort(404)
        get_visible_building(record.building_id, write=True)
        return _service_form(app, record, "telecom", False)

    @app.post("/telecom-usage/<int:record_id>/delete")
    @login_required
    @module_access_required("telecom")
    @write_access_required
    def telecom_usage_delete(record_id):
        record = db.session.get(TelecomUsageRecord, record_id)
        if not record:
            abort(404)
        get_visible_building(record.building_id, write=True)
        db.session.delete(record)
        db.session.commit()
        return redirect(url_for("telecom_usage_list"))

    # ---------------- Reports ----------------
    @app.get("/reports")
    @login_required
    @module_access_required("reports")
    def reports():
        return render_template("reports.html", default_year=_reporting_year())

    @app.get("/reports/export/<report_key>")
    @login_required
    @module_access_required("reports")
    def report_export(report_key):
        if report_key in MODULE_LABELS and not module_allowed(report_key):
            abort(403)
        if report_key in UTILITY_LABELS:
            records = visible_bills_query().filter(UtilityBill.utility_type == report_key).order_by(
                UtilityBill.year.desc(), Building.name, UtilityBill.reading_from
            ).all()
            detail_headers = [
                "شهر", "کد شهر", "ساختمان", "کاربری", "آدرس", "تعداد افراد",
                "شماره اشتراک", "شناسه قبض", "نوع دوره", "سال شمسی", "از تاریخ شمسی",
                "تا تاریخ شمسی", "میزان مصرف", "واحد", "مبلغ (ریال)", "سرانه مصرف",
            ]
            detail_rows = []
            for record in records:
                sub = record.utility_subscription or record.building.subscription(report_key)
                detail_rows.append([
                    record.building.city.name, record.building.city.code, record.building.name,
                    record.building.usage_type.name, record.building.address, record.building.staff_count,
                    sub.subscription_number if sub else None, sub.bill_identifier if sub else None,
                    PERIOD_LABELS[record.period_type], record.year, record.reading_from, record.reading_to,
                    record.consumption, record.unit, _rounded_integer(record.amount), record.per_capita,
                ])
            summary_headers = [
                "شهر", "کد شهر", "ساختمان", "کاربری", "آدرس", "شماره اشتراک", "شناسه قبض",
                "سال شمسی", "مبنای محاسبه", "جمع مصرف", "واحد", "جمع مبلغ (ریال)", "تعداد افراد", "سرانه مصرف",
            ]
            summary_rows = [[
                row["city_name"], row["city_code"], row["building_name"], row["usage_name"], row["address"],
                row["subscription_number"], row["bill_identifier"], row["year"], PERIOD_LABELS[row["source"]],
                _rounded_integer(row["consumption"]), UTILITY_UNITS[row["utility_type"]],
                _rounded_integer(row["amount"]), row["staff_count"], row["per_capita"],
            ] for row in _utility_annual_summaries(utility=report_key)]
            return _multi_sheet_response([
                {"title": "جزئیات دوره‌ها", "headers": detail_headers, "rows": detail_rows,
                 "integer_columns": (6,), "currency_columns": (15,)},
                {"title": "خلاصه سالیانه", "headers": summary_headers, "rows": summary_rows,
                 "integer_columns": (10, 13), "currency_columns": (12,)},
            ], f"{report_key}_report_{_jalali_stamp()}.xlsx")

        if report_key == "production":
            records = visible_production_query().order_by(EnergyProductionRecord.generator_type, EnergyProductionRecord.year.desc(), Building.name).all()
            headers = ["شهر", "کد شهر", "ساختمان", "کاربری", "آدرس", "سال شمسی", "از تاریخ شمسی", "تا تاریخ شمسی", "نوع ژنراتور", "ظرفیت تولید", "انرژی تولیدشده (kWh)", "توضیحات"]
            rows = [[
                r.building.city.name, r.building.city.code, r.building.name, r.building.usage_type.name,
                r.building.address, r.year, r.period_from, r.period_to,
                GENERATOR_TYPES.get(r.generator_type, "سایر"),
                _rounded_integer(r.capacity_kwh) if r.capacity_kwh is not None else None,
                _rounded_integer(r.produced_kwh), r.notes,
            ] for r in records]
            summary_rows = [[
                item["label"], item["building_count"], _rounded_integer(item["capacity"]),
                item["record_count"], _rounded_integer(item["produced"]),
            ] for item in _production_generator_summary()]
            return _multi_sheet_response([
                {"title": "جزئیات تولید", "headers": headers, "rows": rows, "integer_columns": (10, 11)},
                {"title": "تفکیک نوع ژنراتور", "headers": [
                    "نوع ژنراتور", "تعداد ساختمان", "جمع ظرفیت (kWh)", "تعداد رکورد", "جمع تولید (kWh)",
                ], "rows": summary_rows, "integer_columns": (2, 3, 4, 5)},
            ], f"production_report_{_jalali_stamp()}.xlsx")

        if report_key == "telecom":
            records = visible_telecom_query().order_by(
                TelecomUsageRecord.telecom_type, TelecomUsageRecord.year.desc(), Building.name, TelecomUsageRecord.id
            ).all()
            headers = ["شهر", "کد شهر", "ساختمان", "آدرس", "تعداد خطوط", "نوع سرویس", "نوع مخابرات", "سرویس‌دهنده مخابرات", "شماره پرونده", "سرویس‌دهنده اینترنت", "سرعت اینترنت", "سال شمسی", "از تاریخ شمسی", "تا تاریخ شمسی", "مبلغ (ریال)", "شرح مصرف"]
            rows = []
            for r in records:
                telecom = r.building.telecom_service
                internet = r.building.network_service("internet")
                rows.append([
                    r.building.city.name, r.building.city.code, r.building.name, r.building.address,
                    r.building.phone_line_count, SERVICE_KINDS[r.service_kind],
                    TELECOM_TYPES.get(r.telecom_type, "") if r.telecom_type else None,
                    telecom.provider if telecom else None, telecom.case_number if telecom else None,
                    internet.provider if internet else None, internet.speed_mbps if internet else None,
                    r.year, r.period_from, r.period_to, _rounded_integer(r.amount), r.usage_description,
                ])
            return _workbook_response(headers, rows, "سرویس مخابراتی", f"telecom_report_{_jalali_stamp()}.xlsx", integer_columns=(5,), currency_columns=(15,))
        abort(404)

    @app.get("/reports/dynamic")
    @login_required
    @module_access_required("reports")
    @module_access_required("dynamic_reports")
    def dynamic_report():
        datasets = allowed_dynamic_datasets()
        if not datasets:
            abort(403)
        dataset = request.args.get("dataset", "utility" if "utility" in datasets else datasets[0])
        if dataset not in datasets:
            abort(403)
        try:
            result = _dynamic_detail_result(dataset, request.args, limit=500)
        except ValueError as exc:
            flash(str(exc), "danger")
            result = {"columns": [], "rows": [], "selected": [], "truncated": False, "count": 0}
        fields = [{"key": key, "label": spec[0]} for key, spec in _authorized_dynamic_field_specs(dataset).items()]
        cities = City.query.order_by(City.name).all() if current_user.can_view_all_cities else ([current_user.assigned_city] if current_user.assigned_city else [])
        buildings = visible_buildings_query().order_by(Building.name).all()
        query_string = request.query_string.decode("utf-8")
        export_url = url_for("dynamic_report_export") + (f"?{query_string}" if query_string else f"?dataset={dataset}")
        return render_template(
            "dynamic_report.html", dataset=dataset, datasets=datasets, result=result, fields=fields,
            cities=cities, buildings=buildings, export_url=export_url,
        )

    @app.get("/reports/dynamic/export")
    @login_required
    @module_access_required("reports")
    @module_access_required("dynamic_reports")
    def dynamic_report_export():
        datasets = allowed_dynamic_datasets()
        dataset = request.args.get("dataset", "utility" if "utility" in datasets else (datasets[0] if datasets else ""))
        if dataset not in datasets:
            abort(403)
        try:
            result = _dynamic_detail_result(dataset, request.args, limit=10000, excel=True)
        except ValueError as exc:
            flash(str(exc), "danger")
            return redirect(url_for("dynamic_report", dataset=dataset))
        integer_columns = tuple(i + 1 for i, col in enumerate(result["columns"]) if col["kind"] == "integer")
        currency_columns = tuple(i + 1 for i, col in enumerate(result["columns"]) if col["kind"] == "rial")
        return _workbook_response(
            [col["label"] for col in result["columns"]], result["rows"], "گزارش پویا",
            f"dynamic_{dataset}_{_jalali_stamp()}.xlsx", integer_columns=integer_columns,
            currency_columns=currency_columns,
        )

    @app.get("/api/reports/data")
    @login_required
    @module_access_required("reports")
    def reports_data():
        report_year = request.args.get("year", type=int) or _reporting_year()
        comparison_period = request.args.get("period", "periodic")
        if comparison_period not in PERIOD_LABELS:
            comparison_period = "periodic"
        # First aggregate per building/year. Periodic rows take precedence over the
        # optional annual row, preventing double counting in reports.
        q = db.session.query(
            City.name.label("city_name"), UtilityBill.building_id, UtilityBill.utility_type, UtilityBill.year,
            func.sum(case((UtilityBill.period_type == "periodic", UtilityBill.consumption), else_=0)).label("p_cons"),
            func.sum(case((UtilityBill.period_type == "periodic", UtilityBill.amount), else_=0)).label("p_amount"),
            func.sum(case((UtilityBill.period_type == "annual", UtilityBill.consumption), else_=0)).label("a_cons"),
            func.sum(case((UtilityBill.period_type == "annual", UtilityBill.amount), else_=0)).label("a_amount"),
            func.sum(case((UtilityBill.period_type == "periodic", 1), else_=0)).label("p_count"),
        ).join(Building, UtilityBill.building_id == Building.id).join(City, Building.city_id == City.id).filter(
            UtilityBill.utility_type.in_(allowed_utilities())
        )
        if not current_user.can_view_all_cities:
            q = q.filter(Building.city_id == current_user.city_id)
        rows = q.group_by(City.name, UtilityBill.building_id, UtilityBill.utility_type, UtilityBill.year).all()
        cities = sorted({r.city_name for r in rows})
        consumption = {u: {city: 0 for city in cities} for u in UTILITY_LABELS}
        amounts = {u: {city: 0 for city in cities} for u in UTILITY_LABELS}
        for row in rows:
            use_periodic = row.p_count > 0
            consumption[row.utility_type][row.city_name] += _rounded_integer((row.p_cons if use_periodic else row.a_cons) or 0)
            amounts[row.utility_type][row.city_name] += _rounded_integer((row.p_amount if use_periodic else row.a_amount) or 0)
        production_summary = []
        if module_allowed("production"):
            production_summary = [{
                "key": item["generator_type"], "label": item["label"],
                "building_count": item["building_count"], "record_count": item["record_count"],
                "capacity": _rounded_integer(item["capacity"]), "produced": _rounded_integer(item["produced"]),
            } for item in _production_generator_summary()]
        return jsonify({
            "cities": cities, "consumption": consumption, "amounts": amounts,
            "comparison_year": report_year, "comparison_period": comparison_period,
            "per_capita": _city_per_capita_comparison(comparison_period, report_year),
            "production_by_generator": production_summary,
        })

    # ---------------- Normalized city reference administration ----------------
    @app.route("/admin/cities", methods=["GET", "POST"])
    @login_required
    @admin_required
    def admin_cities():
        if request.method == "POST":
            try:
                code = _text(request.form.get("code"), "کد شهر", required=True, max_length=30)
                name = _text(request.form.get("name"), "نام شهر", required=True, max_length=100)
                if City.query.filter(or_(City.code == code, City.name == name)).first():
                    raise ValueError("کد یا نام شهر قبلاً ثبت شده است.")
                db.session.add(City(code=code, name=name, is_active=True))
                db.session.commit()
                flash("شهر جدید ثبت شد.", "success")
                return redirect(url_for("admin_cities"))
            except ValueError as exc:
                db.session.rollback()
                flash(str(exc), "danger")
        cities = City.query.order_by(City.name).all()
        return render_template("admin_cities.html", cities=cities)

    @app.route("/admin/cities/<int:city_id>/edit", methods=["GET", "POST"])
    @login_required
    @admin_required
    def admin_city_edit(city_id):
        city = db.session.get(City, city_id)
        if not city:
            abort(404)
        if request.method == "POST":
            try:
                code = _text(request.form.get("code"), "کد شهر", required=True, max_length=30)
                name = _text(request.form.get("name"), "نام شهر", required=True, max_length=100)
                duplicate = City.query.filter(or_(City.code == code, City.name == name), City.id != city.id).first()
                if duplicate:
                    raise ValueError("کد یا نام شهر قبلاً ثبت شده است.")
                city.code, city.name = code, name
                city.is_active = _bool(request.form.get("is_active"), False)
                db.session.commit()
                flash("مرجع شهر به‌روزرسانی شد و تغییر برای همه ساختمان‌ها اعمال گردید.", "success")
                return redirect(url_for("admin_cities"))
            except ValueError as exc:
                db.session.rollback()
                flash(str(exc), "danger")
        return render_template("admin_city_form.html", city=city)

    # ---------------- User administration ----------------
    @app.get("/admin/users")
    @login_required
    @admin_required
    def admin_users():
        sort = request.args.get("sort", "username")
        direction = request.args.get("dir", "asc")
        column = {"username": User.username, "full_name": User.full_name, "role": User.role, "city": City.name, "active_until": User.active_until, "status": User.is_active_flag}.get(sort, User.username)
        query = User.query.outerjoin(City).order_by(column.asc() if direction == "asc" else column.desc())
        return render_template("admin_users.html", users=query.all(), current_sort=sort, current_dir=direction)

    @app.route("/admin/users/new", methods=["GET", "POST"])
    @login_required
    @admin_required
    def admin_user_new():
        user = User()
        if request.method == "POST" and _apply_user_form(user, is_new=True):
            try:
                db.session.add(user)
                db.session.commit()
                flash("کاربر جدید ایجاد شد.", "success")
                return redirect(url_for("admin_users"))
            except IntegrityError:
                db.session.rollback()
                flash("نام کاربری تکراری است یا ارتباط شهر انتخاب‌شده معتبر نیست.", "danger")
            except Exception:
                db.session.rollback()
                app.logger.exception("admin user create failed")
                flash("ذخیره کاربر ناموفق بود. تاریخ‌ها و اطلاعات واردشده را بررسی کنید.", "danger")
        return render_template("admin_user_form.html", user=None, cities=City.query.order_by(City.name).all())

    @app.route("/admin/users/<int:user_id>/edit", methods=["GET", "POST"])
    @login_required
    @admin_required
    def admin_user_edit(user_id):
        user = db.session.get(User, user_id)
        if not user:
            abort(404)
        if request.method == "POST" and _apply_user_form(user, is_new=False):
            try:
                db.session.commit()
                if user.id == current_user.id and not user.is_active:
                    logout_user()
                    session.clear()
                    return redirect(url_for("login"))
                flash("اطلاعات کاربر به‌روزرسانی شد.", "success")
                return redirect(url_for("admin_users"))
            except IntegrityError:
                db.session.rollback()
                flash("اطلاعات کاربر با محدودیت‌های پایگاه‌داده سازگار نیست.", "danger")
            except Exception:
                db.session.rollback()
                app.logger.exception("admin user update failed")
                flash("ویرایش کاربر ناموفق بود. تاریخ‌ها و اطلاعات واردشده را بررسی کنید.", "danger")
        return render_template("admin_user_form.html", user=user, cities=City.query.order_by(City.name).all())

    @app.post("/admin/users/<int:user_id>/delete")
    @login_required
    @admin_required
    def admin_user_delete(user_id):
        if user_id == current_user.id:
            flash("نمی‌توانید حساب خودتان را حذف کنید.", "danger")
            return redirect(url_for("admin_users"))
        user = db.session.get(User, user_id)
        if user:
            db.session.delete(user)
            db.session.commit()
        return redirect(url_for("admin_users"))

    @app.get("/admin/users/export.xlsx")
    @login_required
    @admin_required
    def admin_users_export_excel():
        headers = ["نام کاربری", "نام کامل", "تلفن", "نقش", "کد شهر", "نام شهر", "فعال", "فعال از", "فعال تا", "آخرین ورود"] + list(MODULE_LABELS.values())
        rows = [[
            u.username, u.full_name, u.phone, u.role,
            u.assigned_city.code if u.assigned_city else None, u.city_name,
            "بله" if u.is_active_flag else "خیر", _jalali_datetime_display(u.active_from),
            _jalali_datetime_display(u.active_until), _jalali_datetime_display(u.last_login_at),
            *["بله" if u.has_module(key) else "خیر" for key in MODULE_LABELS],
        ] for u in User.query.order_by(User.username).all()]
        return _workbook_response(headers, rows, "کاربران", f"users_{_jalali_stamp()}.xlsx")

    @app.get("/admin/users/import-template")
    @login_required
    @admin_required
    def admin_user_import_template():
        headers = ["username", "password", "full_name", "phone", "role", "city_code", "is_active", "active_from", "active_until"] + [f"module_{key}" for key in MODULE_LABELS]
        row = ["user_example", "StrongPass123!", "کاربر نمونه", "09120000000", "user", "091", "بله", "", "۱۴۰۵/۱۲/۲۹ ۲۳:۵۹"] + ["بله" for _ in MODULE_LABELS]
        return _workbook_response(headers, [row], "کاربران", "user_import_template.xlsx")

    @app.route("/admin/users/import", methods=["GET", "POST"])
    @login_required
    @admin_required
    def admin_user_import():
        result = None
        if request.method == "POST":
            upload = request.files.get("excel_file")
            if not upload or not upload.filename or Path(upload.filename).suffix.lower() not in (".xlsx", ".xlsm"):
                flash("فایل xlsx یا xlsm معتبر انتخاب کنید.", "danger")
                return redirect(url_for("admin_user_import"))
            try:
                rows = list(load_workbook(upload, read_only=True, data_only=True).active.iter_rows(values_only=True))
                if not rows:
                    raise ValueError("فایل خالی است.")
                headers = [str(x).strip() if x else "" for x in rows[0]]
                created, skipped, errors = [], [], []
                for idx, row in enumerate(rows[1:], 2):
                    if not row or all(x is None or str(x).strip() == "" for x in row):
                        continue
                    data = {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
                    username = str(data.get("username") or "").strip()
                    if User.query.filter(func.lower(User.username) == username.casefold()).first():
                        skipped.append(f"ردیف {idx}: {username}")
                        continue
                    user = User()
                    try:
                        if _apply_user_mapping(user, data, is_new=True):
                            db.session.add(user)
                            created.append(username)
                    except ValueError as exc:
                        errors.append(f"ردیف {idx}: {exc}")
                db.session.commit()
                result = {"created": created, "skipped": skipped, "errors": errors}
            except Exception as exc:
                db.session.rollback()
                app.logger.exception("user import failed")
                flash(f"پردازش فایل ناموفق بود: {exc}", "danger")
        return render_template("admin_user_import.html", result=result)

    @app.get("/admin/users/export")
    @login_required
    @admin_required
    def admin_users_export():
        users = User.query.order_by(User.id).all()
        payload = {
            "version": 3,
            "exported_at": utcnow().isoformat(),
            "users": [{
                "username": u.username, "password_hash": u.password_hash, "full_name": u.full_name,
                "phone": u.phone, "role": u.role, "city_code": u.assigned_city.code if u.assigned_city else None,
                "is_active": u.is_active_flag, "active_from": (u.active_from.isoformat() + "Z") if u.active_from else None,
                "active_until": (u.active_until.isoformat() + "Z") if u.active_until else None,
                "permissions": {key: u.has_module(key) for key in MODULE_LABELS},
            } for u in users],
        }
        stream = io.BytesIO(json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8"))
        return send_file(stream, as_attachment=True, download_name=f"users_backup_{utcnow():%Y%m%d_%H%M%S}.json", mimetype="application/json")

    @app.route("/admin/users/restore", methods=["GET", "POST"])
    @login_required
    @admin_required
    def admin_users_restore():
        result = None
        if request.method == "POST":
            upload = request.files.get("backup_file")
            try:
                if not upload or not upload.filename:
                    raise ValueError("فایل پشتیبان انتخاب نشده است.")
                payload = json.load(upload)
                if not isinstance(payload.get("users"), list):
                    raise ValueError("ساختار فایل پشتیبان معتبر نیست.")
                mode = request.form.get("mode", "skip")
                created, updated, skipped = [], [], []
                for item in payload["users"]:
                    username = _text(item.get("username"), "نام کاربری", required=True, max_length=80)
                    existing = User.query.filter_by(username=username).first()
                    if existing and mode != "overwrite":
                        skipped.append(username)
                        continue
                    user = existing or User(username=username)
                    user.full_name = _text(item.get("full_name"), "نام کامل", required=True, max_length=150)
                    user.phone = _text(item.get("phone"), "تلفن", max_length=20)
                    user.role = item.get("role") if item.get("role") in ROLES else "user"
                    city_code = item.get("city_code")
                    user.assigned_city = City.query.filter_by(code=str(city_code)).first() if city_code else None
                    user.is_active_flag = _bool(item.get("is_active"), True)
                    user.active_from = _parse_datetime_local(item.get("active_from"), "فعال از")
                    user.active_until = _parse_datetime_local(item.get("active_until"), "فعال تا")
                    password_hash = item.get("password_hash")
                    if not isinstance(password_hash, str) or not password_hash:
                        raise ValueError(f"هش رمز کاربر {username} نامعتبر است.")
                    user.password_hash = password_hash
                    permissions = item.get("permissions") if isinstance(item.get("permissions"), dict) else {}
                    permission_rows = {row.module_key: row for row in user.module_permissions}
                    for module_key in MODULE_LABELS:
                        enabled = _bool(permissions.get(module_key), True)
                        permission = permission_rows.get(module_key)
                        if not permission:
                            permission = UserModulePermission(user=user, module_key=module_key)
                            db.session.add(permission)
                        permission.is_enabled = enabled
                    if not existing:
                        db.session.add(user)
                        created.append(username)
                    else:
                        updated.append(username)
                db.session.commit()
                result = {"created": created, "updated": updated, "skipped": skipped}
            except (ValueError, KeyError, json.JSONDecodeError) as exc:
                db.session.rollback()
                flash(str(exc), "danger")
        return render_template("admin_users_restore.html", result=result)


def _bill_form(app, bill, is_new):
    if request.method == "POST":
        try:
            building_id = _integer(request.form.get("building_id"), "ساختمان", required=True, minimum=1)
            building = get_visible_building(building_id, write=True)
            utility = request.form.get("utility_type")
            period_type = request.form.get("period_type")
            if utility not in UTILITY_LABELS or not module_allowed(utility):
                raise ValueError("نوع انشعاب نامعتبر یا خارج از سطح دسترسی است.")
            subscription_id = _integer(
                request.form.get("utility_subscription_id"), "اشتراک", required=True, minimum=1
            )
            subscription = db.session.get(UtilitySubscription, subscription_id)
            if not subscription or subscription.building_id != building.id or subscription.utility_type != utility:
                raise ValueError("اشتراک انتخاب‌شده متعلق به ساختمان و انشعاب انتخابی نیست.")
            if period_type not in PERIOD_LABELS:
                raise ValueError("نوع دوره نامعتبر است.")
            year = _integer(request.form.get("year"), "سال", required=True, minimum=1300)
            reading_from = _parse_date(request.form.get("reading_from"), "تاریخ شروع", required=True)
            reading_to = _parse_date(request.form.get("reading_to"), "تاریخ پایان", required=True)
            if reading_to < reading_from:
                raise ValueError("تاریخ پایان نمی‌تواند قبل از تاریخ شروع باشد.")
            duplicate = UtilityBill.query.filter_by(
                utility_subscription_id=subscription.id, period_type=period_type, year=year,
                reading_from=reading_from, reading_to=reading_to,
            )
            if not is_new:
                duplicate = duplicate.filter(UtilityBill.id != bill.id)
            if duplicate.first():
                raise ValueError("برای این ساختمان، انشعاب و بازه زمانی قبلاً قبض ثبت شده است.")
            if period_type == "annual":
                annual = UtilityBill.query.filter_by(utility_subscription_id=subscription.id, period_type="annual", year=year)
                if not is_new:
                    annual = annual.filter(UtilityBill.id != bill.id)
                if annual.first():
                    raise ValueError("برای هر ساختمان و انشعاب در هر سال فقط یک رکورد سالیانه مجاز است.")
            bill.building = building
            bill.utility_subscription = subscription
            bill.utility_type = utility
            bill.period_type = period_type
            bill.year = year
            bill.reading_from = reading_from
            bill.reading_to = reading_to
            bill.consumption = _decimal(request.form.get("consumption"), "میزان مصرف", required=True)
            bill.amount = _integer(request.form.get("amount"), "مبلغ به ریال", required=True)
            bill.notes = _text(request.form.get("notes"), "توضیحات", max_length=2000)
            if is_new:
                db.session.add(bill)
            db.session.commit()
            flash("قبض و مصرف با موفقیت ذخیره شد.", "success")
            return redirect(url_for("consumptions", utility=utility, year=year))
        except ValueError as exc:
            db.session.rollback()
            flash(str(exc), "danger")
        except Exception:
            db.session.rollback()
            app.logger.exception("bill save failed")
            flash("ذخیره قبض ناموفق بود.", "danger")
    utilities = allowed_utilities()
    if not utilities:
        abort(403)
    buildings = visible_buildings_query().filter_by(is_active=True).options(
        selectinload(Building.utility_subscriptions)
    ).order_by(Building.name).all()
    return render_template(
        "consumption_form.html", bill=None if is_new else bill, buildings=buildings,
        default_year=_reporting_year(), default_utility=utilities[0],
    )


def _service_form(app, record, kind, is_new):
    if request.method == "POST":
        try:
            building_id = _integer(request.form.get("building_id"), "ساختمان", required=True, minimum=1)
            building = get_visible_building(building_id, write=True)
            record.building = building
            record.year = _integer(request.form.get("year"), "سال", required=True, minimum=1300)
            record.period_from = _parse_date(request.form.get("period_from"), "تاریخ شروع", required=True)
            record.period_to = _parse_date(request.form.get("period_to"), "تاریخ پایان", required=True)
            if record.period_to < record.period_from:
                raise ValueError("تاریخ پایان قبل از تاریخ شروع است.")
            if kind == "production":
                generator_type = request.form.get("generator_type")
                if generator_type not in GENERATOR_TYPES:
                    raise ValueError("نوع ژنراتور نامعتبر است.")
                generator = _get_or_create_related(building, "generator", PowerGenerator)
                generator.generator_type = generator_type
                generator.production_capacity_kwh = _decimal(
                    request.form.get("production_capacity_kwh"), "ظرفیت تولید", required=True
                )
                record.generator_type = generator_type
                record.capacity_kwh = generator.production_capacity_kwh
                record.produced_kwh = _decimal(request.form.get("produced_kwh"), "انرژی تولیدشده", required=True)
                record.notes = _text(request.form.get("notes"), "توضیحات", max_length=2000)
                endpoint = "production_list"
            else:
                service_kind = request.form.get("service_kind")
                if service_kind not in SERVICE_KINDS:
                    raise ValueError("نوع سرویس نامعتبر است.")
                telecom_type = request.form.get("telecom_type")
                if telecom_type not in TELECOM_TYPES:
                    raise ValueError("نوع سرویس مخابراتی نامعتبر است.")
                record.service_kind = service_kind
                record.telecom_type = telecom_type
                record.amount = _integer(request.form.get("amount"), "مبلغ به ریال", required=True)
                record.usage_description = _text(request.form.get("usage_description"), "شرح مصرف", max_length=500)
                endpoint = "telecom_usage_list"
            if is_new:
                db.session.add(record)
            db.session.commit()
            flash("رکورد ذخیره شد.", "success")
            return redirect(url_for(endpoint))
        except ValueError as exc:
            db.session.rollback()
            flash(str(exc), "danger")
        except Exception:
            db.session.rollback()
            app.logger.exception("service record save failed")
            flash("ذخیره رکورد ناموفق بود.", "danger")
    buildings = visible_buildings_query().filter_by(is_active=True).order_by(Building.name).all()
    return render_template("service_record_form.html", record=None if is_new else record, kind=kind, buildings=buildings, default_year=_reporting_year())


def _apply_user_mapping(user, source, is_new):
    username = _text(source.get("username"), "نام کاربری", required=is_new, max_length=80)
    if is_new:
        if not re.fullmatch(r"[A-Za-z0-9_.-]{3,80}", username or ""):
            raise ValueError("نام کاربری باید ۳ تا ۸۰ نویسه و فقط شامل حروف لاتین، عدد، نقطه، خط تیره یا زیرخط باشد.")
        if User.query.filter(func.lower(User.username) == username.casefold()).first():
            raise ValueError("این نام کاربری قبلاً استفاده شده است.")
        user.username = username
    user.full_name = _text(source.get("full_name"), "نام کامل", required=True, max_length=150)
    user.phone = _text(source.get("phone"), "تلفن", max_length=20)
    role = source.get("role") or "user"
    if role not in ROLES:
        raise ValueError("نقش کاربر نامعتبر است.")
    user.role = role
    city_id = source.get("city_id")
    city_code = source.get("city_code")
    city = db.session.get(City, int(city_id)) if city_id and str(city_id).isdigit() else None
    if not city and city_code:
        city = City.query.filter_by(code=str(city_code).strip()).first()
    if role in ("user", "view") and not city:
        raise ValueError("برای کاربر شهرستانی، انتخاب شهر الزامی است.")
    user.assigned_city = city if role in ("user", "view") else None
    user.is_active_flag = _bool(source.get("is_active"), default=True)
    if source.get("active_from_date") or source.get("active_until_date"):
        user.active_from = _parse_jalali_datetime(
            source.get("active_from_date"), source.get("active_from_time"), "شروع فعالیت", "00:00"
        )
        user.active_until = _parse_jalali_datetime(
            source.get("active_until_date"), source.get("active_until_time"), "پایان فعالیت", "23:59"
        )
    else:
        user.active_from = _parse_datetime_local(source.get("active_from"), "فعال از")
        user.active_until = _parse_datetime_local(source.get("active_until"), "فعال تا")
    if user.active_from and user.active_until and user.active_until <= user.active_from:
        raise ValueError("تاریخ پایان فعالیت باید بعد از تاریخ شروع باشد.")
    password = str(source.get("password") or "")
    if is_new or password:
        if len(password) < 10 or not re.search(r"[A-Za-z]", password) or not re.search(r"\d", password):
            raise ValueError("رمز عبور باید حداقل ۱۰ نویسه و شامل حرف و عدد باشد.")
        user.set_password(password)

    existing_permissions = {item.module_key: item for item in user.module_permissions}
    permission_payload = source.get("permissions") if hasattr(source, "get") else None
    for module_key in MODULE_LABELS:
        raw = source.get(f"module_{module_key}")
        if isinstance(permission_payload, dict) and module_key in permission_payload:
            raw = permission_payload[module_key]
        if raw is None and not is_new:
            continue
        enabled = _bool(raw, default=True)
        permission = existing_permissions.get(module_key)
        if not permission:
            permission = UserModulePermission(user=user, module_key=module_key)
            db.session.add(permission)
        permission.is_enabled = enabled
    return True


def _apply_user_form(user, is_new):
    try:
        return _apply_user_mapping(user, request.form, is_new)
    except ValueError as exc:
        db.session.rollback()
        flash(str(exc), "danger")
        return False
    except Exception:
        db.session.rollback()
        current_app.logger.exception("user form validation failed")
        flash("اعتبارسنجی کاربر ناموفق بود. قالب تاریخ و شهر انتخاب‌شده را بررسی کنید.", "danger")
        return False


app = create_app()

if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5700"))
    app.run(host="0.0.0.0", port=port, debug=os.environ.get("FLASK_DEBUG") == "1")
